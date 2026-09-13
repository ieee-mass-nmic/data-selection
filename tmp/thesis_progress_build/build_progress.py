from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path('/Users/bytedance/code/data-selection')
OUT = ROOT / 'output/thesis/学位论文进展情况及下一阶段工作重点.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.70), Inches(.68)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)
sec.footer_distance = Inches(.30)

def font_style(style, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), 'Arial Unicode MS')
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')

font_style(doc.styles['Normal'], 11.5)
n = doc.styles['Normal'].paragraph_format
n.line_spacing = 1.17
n.space_after = Pt(6)
n.widow_control = True
for sn, size in [('Title', 18), ('Subtitle', 10.5), ('Heading 1', 14), ('Heading 2', 12)]:
    font_style(doc.styles[sn], size, sn != 'Subtitle')
    pf = doc.styles[sn].paragraph_format
    pf.space_before = Pt(10 if sn.startswith('Heading') else 0)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    pf.line_spacing = 1.15
    if sn.startswith('Heading'):
        doc.styles[sn].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
doc.styles['Subtitle'].font.italic = False

def p(text, bold_prefix=None, size=None):
    para = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold = True
        para.add_run(text[len(bold_prefix):])
    else:
        para.add_run(text)
    if size:
        for r in para.runs: r.font.size = Pt(size)
    return para

def heading(text, page=False):
    para = doc.add_paragraph(text, 'Heading 1')
    if page: para.paragraph_format.page_break_before = True
    return para

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for c, h, width in zip(t.rows[0].cells, headers, widths):
        c.width = Inches(width)
        c.text = h
    for row in rows:
        cells = t.add_row().cells
        for c, text, width in zip(cells, row, widths):
            c.width = Inches(width)
            c.text = text
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement('w:'+edge)
        b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9')
        borders.append(b)
    pr.append(borders)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                item = OxmlElement('w:'+edge); item.set(qn('w:w'),'80'); item.set(qn('w:type'),'dxa'); margins.append(item)
            tcpr.append(margins)
            if ri == 0:
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'),'E8EEF3'); tcpr.append(shd)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.08
                for run in para.runs:
                    run.font.size = Pt(10.5)
                    if ri == 0: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def ref(label, url):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.05
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    s = OxmlElement('w:sz'); s.set(qn('w:val'),'19'); rp.append(s)
    color = OxmlElement('w:color'); color.set(qn('w:val'),'244A64'); rp.append(color)
    r.append(rp); tx = OxmlElement('w:t'); tx.text = label; r.append(tx); hyperlink.append(r)
    para._p.append(hyperlink)
    return para

doc.add_paragraph('学位论文进展情况及下一阶段工作重点', 'Title')
doc.add_paragraph('2026年9月7日', 'Subtitle')
p('本学位论文围绕大语言模型参数高效微调中的数据选择问题展开。目前已完成第一项工作 PCU-Select 的论文撰写，下一阶段拟以已有算法和代码为基础，开展多配置数据选择与微调实验管理系统的设计、实现和验证，形成算法研究与工程应用相衔接的两个工作点。')

heading('一 当前学位论文进展情况')
p('研究问题与方法。第一项工作针对不同参数高效微调（PEFT）配置下样本价值不同、逐配置重建梯度数据存储开销较高的问题，提出 PEFT 条件效用建模方法。通过24个共享干预站点、样本与任务表征以及结构化配置编码，复用离线计算，为不同任务和配置生成各自的数据子集。', '研究问题与方法。')
p('实现与实验基础。已形成特征缓存、多保真标签构造、条件评分、聚类配额选择及配置迁移校准的实现框架。论文稿已呈现 Llama-2-7B 上 GSM8K、HumanEval、MMLU 和 TyDiQA 四项任务、五种已见 PEFT 配置的主要比较，并包含动机实验、消融、迁移和成本分析。', '实现与实验基础。')
p('阶段性结果。论文当前报告的四任务指标宏平均为35.18，逐配置 LESS 为35.14，支持平均效果接近的结论。在一次性离线投入144.0 GPU小时后，支持范围内新增一个四任务 PEFT 配置的边际选数成本由63.2降至1.51 GPU小时，约降至原来的1/42；该比较不包含下游微调成本，也不表示训练全过程加速42倍。', '阶段性结果。')
p('当前阶段与待完善内容。现有实现主要通过脚本组织实验，已有离线缓存与配置注册基础，但尚需完善产物版本依赖、批量请求复用、结果隔离和失败恢复。后续将核对原始日志与图表口径，把第一项工作整理为算法章节，并将上述工程问题作为第二项工作的切入点。', '当前阶段与待完善内容。')

heading('二 下一阶段工作重点')
p('下一阶段以 PCU-Select 的工程化落地为重点，将已有特征提取、评分、聚类选数及训练评测模块封装为统一流程。系统支持登记数据池和任务、批量提交 PEFT 配置与选数预算、查看执行状态，以及导出所选数据、指标和成本报告。')
p('核心工作是实现依赖感知的分层缓存、增量复用和公共计算去重。通过重复请求、数据追加和任务中断实验，验证结果一致性、执行效率与恢复能力，完成系统原型、实验分析及学位论文撰写。')

heading('三 第二个工作点简要方案', page=True)
p('建议题目  基于 PCU-Select 的多配置数据选择与微调实验管理系统', size=12)
p('第二项工作直接使用第一项工作的评分器、特征缓存、PEFT 编码和聚类配额算法，解决重复开展多任务、多配置实验时的产物管理与执行效率问题。预期成果是一套可运行系统，以及两项能够定量验证的工程机制。')
table(['已有基础','第二项工作增加的内容'],[
    ['样本特征和任务表征缓存','记录内容及版本依赖，识别可复用产物并局部重算'],
    ['条件评分与聚类配额选数','复用公共评分和聚类，为各预算独立生成子集'],
    ['PEFT 注册及训练评测脚本','统一批量执行、阶段恢复、结果隔离与报告导出'],
], [2.45,4.49])
p('工程机制一 分层缓存与增量复用。为样本特征、任务表示、评分、聚类和选择结果建立依赖清单，记录样本内容摘要、模型与评分器版本、预处理配置、随机种子及产物校验值。输入变化后，沿依赖关系仅重算受影响的节点，避免同名样本或不同预算错误命中旧结果。', '工程机制一 分层缓存与增量复用。')
p('复用规则。仅改变选数预算时，复用评分和聚类，重算配额与子集；选择种子不变、仅改变下游训练种子时，复用选择结果。增加候选样本且评分条件不变时，复用旧样本的兼容特征与逐样本评分，只补算新增样本；完整池的聚类与配额仍需重算。候选顺序纳入依赖，首版固定评分器版本。', '复用规则。')
p('工程机制二 公共计算去重与阶段恢复。把实验请求拆分为特征、任务表示、评分、聚类、选数、训练和评测节点；输入依赖相同的节点只执行一次。通过原子写入、完成标记和请求独立目录，避免结果覆盖；中断后复用已完成产物，仅重跑失败阶段。', '工程机制二 公共计算去重与阶段恢复。')
p('具体示例。同一数据池上的2种 PEFT、2个任务、3个预算和2个训练种子，对应24次训练实验。固定特征与评分器版本、聚类配置及选择种子后，可共享4份目标评分和1份聚类，生成12份子集，再完成24次训练。各任务与配置仍使用各自子集；上述次数不等于端到端加速倍数。', '具体示例。')
p('系统范围。首版采用单工作节点、任务队列、元数据索引和本地产物库，提供简洁控制台。操作流程为登记数据与任务、提交实验矩阵、查看复用计划、执行选数、按需训练评测、导出报告。继续沿用第一项工作已验证的骨干与 PEFT 支持范围。', '系统范围。')

heading('四 实验验证与推进安排', page=True)
p('结果正确性。沿用第一项工作的小规模任务与配置，固定算法、硬件、精度和随机种子，对照优化前后的评分数值、所选样本及顺序。覆盖文本变更但样本 ID 不变、预算变化、评分器版本变化等情况，验证缓存正确失效；代表性训练评测用于确认系统优化未改变原算法效果。', '结果正确性。')
p('复用与增量效率。按1、3、5种受支持配置组织多任务、多预算请求，分别比较顺序执行、分层缓存、缓存加公共节点去重。候选池追加1%、5%、10%样本，对照全量重算与增量执行，记录重算样本数、重复节点数、完成时间、缓存命中率和存储开销。', '复用与增量效率。')
p('恢复与对照口径。在产物写入和阶段切换处模拟中断，检查半成品不会被复用，以及恢复后的结果一致性。对照组保留已有特征缓存与模型复用，并消除同一请求内明显的重复评分。冷启动、热缓存和追加场景分别报告，CPU 耗时与实际 GPU 占用分开统计，共享节点仅计费一次。', '恢复与对照口径。')
p('首版与系统验证建议按6至8周推进，优先完成核心执行流程，再补齐展示界面。')
table(['阶段','工作重点','阶段产出'],[
    ['第1周','梳理已有模块，定义产物依赖、缓存边界与基准请求','系统设计及验证方案'],
    ['第2至3周','实现版本缓存、增量特征与评分、公共节点去重','可运行的核心流程'],
    ['第4至5周','接入训练评测，完成阶段恢复、控制台和报告导出','可演示的系统原型'],
    ['第6至8周','完成正确性、性能及恢复实验，整理学位论文章节','系统实验结果与论文稿'],
], [.90,3.90,2.14])
p('论文贡献组织。第一项工作论证条件效用建模和跨配置选择质量；第二项工作论证面向该算法的复用边界、增量执行与系统可靠性。通用缓存和实验追踪已有工具支持[1–2]，本工作的工程贡献应由 PCU-Select 专用依赖设计及系统对照实验支撑，主要评价目标是质量保持、效率改善和结果可追溯。', '论文贡献组织。')

doc.add_paragraph('实现参考', 'Heading 2')
ref('[1] DVC 官方文档 Running Pipelines  依赖检查与运行缓存机制', 'https://github.com/treeverse/dvc.org/blob/main/content/docs/user-guide/pipelines/running-pipelines.md')
ref('[2] MLflow 官方文档 ML Experiment Tracking  实验参数 指标与产物记录', 'https://mlflow.org/docs/latest/ml/tracking/')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run(); fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
doc.core_properties.title = '学位论文进展情况及下一阶段工作重点'
doc.core_properties.subject = 'PCU-Select研究进展与工程化第二工作点方案'
doc.core_properties.author = ''
doc.core_properties.keywords = '学位论文, PCU-Select, PEFT, 系统设计, 缓存复用, 实验管理'
for root in [doc.styles.element, doc.element]:
    for fonts in root.iter(qn('w:rFonts')):
        for key in list(fonts.attrib):
            if 'theme' in key.lower(): del fonts.attrib[key]
        fonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
    for borders in list(root.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
doc.save(OUT)
print(OUT)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables))
