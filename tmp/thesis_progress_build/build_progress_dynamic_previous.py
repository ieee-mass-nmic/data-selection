from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path('/Users/bytedance/code/data-selection')
OUT = ROOT / 'output/thesis/学位论文进展情况及下一阶段工作重点.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.70), Inches(.68)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)
sec.footer_distance = Inches(.30)

def font_style(style, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), 'Arial Unicode MS')
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')

font_style(doc.styles['Normal'], 11.5)
n = doc.styles['Normal'].paragraph_format
n.line_spacing = 1.17
n.space_after = Pt(6)
n.widow_control = True
for sn, size in [('Title', 18), ('Subtitle', 10.5), ('Heading 1', 14), ('Heading 2', 12)]:
    font_style(doc.styles[sn], size, sn != 'Subtitle')
    pf = doc.styles[sn].paragraph_format
    pf.space_before = Pt(10 if sn.startswith('Heading') else 0)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    pf.line_spacing = 1.15
    if sn.startswith('Heading'):
        doc.styles[sn].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
doc.styles['Subtitle'].font.italic = False

def p(text, bold_prefix=None, size=None):
    para = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold = True
        para.add_run(text[len(bold_prefix):])
    else:
        para.add_run(text)
    if size:
        for r in para.runs: r.font.size = Pt(size)
    return para

def heading(text, page=False):
    para = doc.add_paragraph(text, 'Heading 1')
    if page: para.paragraph_format.page_break_before = True
    return para

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for c, h, width in zip(t.rows[0].cells, headers, widths):
        c.width = Inches(width)
        c.text = h
    for row in rows:
        cells = t.add_row().cells
        for c, text, width in zip(cells, row, widths):
            c.width = Inches(width)
            c.text = text
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement('w:'+edge)
        b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9')
        borders.append(b)
    pr.append(borders)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                item = OxmlElement('w:'+edge); item.set(qn('w:w'),'80'); item.set(qn('w:type'),'dxa'); margins.append(item)
            tcpr.append(margins)
            if ri == 0:
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'),'E8EEF3'); tcpr.append(shd)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.08
                for run in para.runs:
                    run.font.size = Pt(10.5)
                    if ri == 0: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def ref(label, url):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.05
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    s = OxmlElement('w:sz'); s.set(qn('w:val'),'19'); rp.append(s)
    color = OxmlElement('w:color'); color.set(qn('w:val'),'244A64'); rp.append(color)
    r.append(rp); tx = OxmlElement('w:t'); tx.text = label; r.append(tx); hyperlink.append(r)
    para._p.append(hyperlink)
    return para

doc.add_paragraph('学位论文进展情况及下一阶段工作重点', 'Title')
doc.add_paragraph('2026年9月7日', 'Subtitle')
p('本学位论文围绕大语言模型参数高效微调中的数据选择问题展开。目前已完成第一项工作 PCU-Select 的论文撰写，下一阶段拟开展训练状态感知的动态数据选择研究，使两个工作点形成从跨配置数据选择到训练过程自适应选择的递进关系。')

heading('一 当前学位论文进展情况')
p('研究问题与方法。第一项工作针对不同参数高效微调（PEFT）配置下样本价值不同、逐配置重建梯度数据存储开销较高的问题，提出 PEFT 条件效用建模方法。通过24个共享干预站点、样本与任务表征以及结构化配置编码，复用离线计算，为不同任务和配置生成各自的数据子集。', '研究问题与方法。')
p('实现与实验基础。已形成特征缓存、多保真标签构造、条件评分、聚类配额选择及配置迁移校准的实现框架。论文稿已呈现 Llama-2-7B 上 GSM8K、HumanEval、MMLU 和 TyDiQA 四项任务、五种已见 PEFT 配置的主要比较，并包含动机实验、消融、迁移和成本分析。', '实现与实验基础。')
p('阶段性结果。论文当前报告的四任务指标宏平均为35.18，逐配置 LESS 为35.14，支持平均效果接近的结论。在一次性离线投入144.0 GPU小时后，支持范围内新增一个四任务 PEFT 配置的边际选数成本由63.2降至1.51 GPU小时，约降至原来的1/42；该比较不包含下游微调成本，也不表示训练全过程加速42倍。', '阶段性结果。')
p('当前阶段与待完善内容。第一项工作已形成论文稿、实验图表和代码基础。现有验证主要覆盖一个骨干模型家族；选数采用训练前固定子集，尚未回答样本价值随训练推进如何变化。后续需继续核对原始日志、统计汇总与图表口径，并将第一项工作整理为学位论文的独立章节。', '当前阶段与待完善内容。')

heading('二 下一阶段工作重点')
p('下一阶段首先验证同一任务与 PEFT 配置下，样本效用排序是否在训练初期、中期和后期发生超过测量噪声的变化。在此基础上，复用 PCU-Select 的静态评分与特征缓存，研究结合少量训练反馈的评分修正和按需刷新机制。')
p('工作重点包括完成动态效用的动机实验、开发可恢复训练状态的探针测量接口、实现轻量状态残差模型，以及开展等训练预算和等总计算预算的对照实验。待收益与开销得到验证后，完善消融分析，并完成第二项工作及学位论文整体撰写。')

heading('三 第二个工作点简要方案', page=True)
p('建议题目  面向参数高效微调的训练状态感知动态数据选择方法', size=12)
p('核心问题是：在任务与 PEFT 配置固定时，如何用较低的更新代价识别已经过时的数据排序，并为后续训练调整采样策略。预期贡献聚焦于训练状态条件下的效用修正与刷新成本控制。')
table(['比较维度','第一项工作 PCU-Select','第二项工作 拟开展'],[
    ['研究对象','不同 PEFT 配置的样本价值差异','同一配置在不同训练状态下的价值变化'],
    ['选择方式','训练前一次性选定数据子集','训练过程中按需调整数据采样'],
    ['主要贡献','共享表征与跨配置计算复用','状态条件残差与低开销刷新机制'],
], [1.0,2.97,2.97])

p('技术路线一 验证排序漂移。先用一个任务、两种 PEFT 配置、三个训练检查点和128个共同探针样本，测量当前状态下短程更新带来的任务反馈集损失变化。比较排序相关性、Top-k 重合率及重复测量噪声，并通过交换阶段子集的短程续训，检查排序变化是否影响后续学习。', '技术路线一 验证排序漂移。')
p('技术路线二 学习状态修正。冻结原评分器，保留静态评分；将反馈集损失变化、探针响应和训练进度编码为状态向量，与样本和 PEFT 编码共同输入轻量残差模型，以同一任务、配置及阶段内归一化的效用排序作为监督，预测排序修正量。先用少量阶段标签验证，再研究跨阶段和配置共享的残差模型。', '技术路线二 学习状态修正。')
p('技术路线三 按需刷新采样。以固定间隔做低成本状态检查，仅在探针排序或残差误差变化超过开发集确定的阈值时，重新评分并调整下一阶段的聚类配额或采样概率。限制刷新次数和探针预算，保留少量均匀探索及上一阶段高价值样本，控制数据更换幅度。', '技术路线三 按需刷新采样。')
p('实现关键。探针试更新必须从当前 PEFT 参数及优化器状态出发，并在测量后恢复参数、优化器、调度器和随机数状态。现有离线短更新接口不能直接代表训练中的真实增量收益；同时应在独立探针上检验固定特征缓存是否仍能支持准确排序。', '实现关键。')
p('独立性与创新边界。第一篇的初始及暖启动锚点用于离线标签构造，并未让评分器随实际训练状态改变决策。GREATS 已研究在线批选择，GAIA 已研究全局在线估值，Filter-then-Weight 已考虑优化器状态[1–3]。因此，本方案需验证“跨 PEFT 可复用的状态修正与按需刷新”是否具有增量价值，不能仅把动态选数本身作为创新。', '独立性与创新边界。')

heading('四 实验验证与推进安排', page=True)
p('最小主实验。先沿用 Llama-2-7B、GSM8K 与 MMLU，以及注意力 LoRA 与 Adapter 两种配置，完成配对的三个随机种子实验。比较静态 PCU-Select、定期随机刷新、基于损失的动态采样、固定周期残差刷新和按需残差刷新，并加入 GREATS 作为已有在线方法对照。探索阶段可缩小候选池；扩大到现有300K池后再报告正式结论。', '最小主实验。')
p('评价与公平性。分别比较相同训练 token 数和相同总 GPU 小时下的任务指标，并记录达到同等性能的成本、累计接触的不同样本数及刷新次数。总成本计入探针、状态编码、评分和训练；共享离线成本同时给出总额及跨配置摊销口径。增加“使用动态方法所见数据并集的静态训练”对照，区分训练时机收益与数据覆盖收益。', '评价与公平性。')
p('必要消融与判定。去掉训练状态输入、去掉 PEFT 条件、改为固定周期刷新，检验各部分贡献。若主张跨配置迁移，必须留出未参与残差训练的 PEFT 配置。反馈集、开发集与测试集保持分离；触发阈值只在开发设置上确定，最终以测试集任务指标评价。', '必要消融与判定。')
p('建议按8周组织首轮研究，实际时长随 GPU 资源和预实验结果调整。')
table(['阶段','工作重点','阶段产出'],[
    ['第1至2周','梳理相关工作，核对第一项工作材料，完成排序漂移及续训验证','研究边界与可行性结论'],
    ['第3至4周','实现训练状态恢复、残差评分和按需刷新，记录完整成本','可运行原型与小规模结果'],
    ['第5至6周','完成主对照、关键消融和预算比较','性能与成本结果及误差分析'],
    ['第7至8周','补充必要验证，整理两项工作关系，撰写论文','第二工作章节与学位论文稿'],
], [.90,3.90,2.14])
p('推进条件。排序变化超过噪声，且动态方法改善性能与总成本的权衡后，再扩大实验。若转向数据与 PEFT 联合优化，则需新增跨配置可比的性能代理，原有组内归一化效用分数不能直接用于配置优选。', '推进条件。')

doc.add_paragraph('参考文献', 'Heading 2')
ref('[1] Wang et al. GREATS: Online Selection of High-Quality Data for LLM Training in Every Iteration. NeurIPS 2024.', 'https://proceedings.neurips.cc/paper_files/paper/2024/hash/ed165f2ff227cf36c7e3ef88957dadd9-Abstract-Conference.html')
ref('[2] Wang et al. Online Data Selection for Instruction Tuning via Gaussian Processes. arXiv:2606.30077, 2026 预印本.', 'https://arxiv.org/abs/2606.30077')
ref('[3] Wang et al. Filter-then-Weight: Online Data Selection and Reweighting for LLM Fine-Tuning. arXiv:2604.00001 预印本.', 'https://arxiv.org/abs/2604.00001')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run(); fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
doc.core_properties.title = '学位论文进展情况及下一阶段工作重点'
doc.core_properties.subject = 'PCU-Select研究进展与第二工作点方案'
doc.core_properties.author = ''
doc.core_properties.keywords = '学位论文, PCU-Select, PEFT, 动态数据选择'
for root in [doc.styles.element, doc.element]:
    for fonts in root.iter(qn('w:rFonts')):
        for key in list(fonts.attrib):
            if 'theme' in key.lower(): del fonts.attrib[key]
        fonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
    for borders in list(root.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
doc.save(OUT)
print(OUT)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables))
