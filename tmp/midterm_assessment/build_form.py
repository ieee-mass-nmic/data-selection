from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
import json, hashlib, re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/Users/bytedance/code/data-selection')
TMP=ROOT/'tmp/midterm_assessment'
OUT=ROOT/'output/midterm_assessment'; OUT.mkdir(parents=True,exist_ok=True)
REF=Path('/Users/bytedance/Desktop/师兄的文档参考/陈瑾如-中期考核表.docx')
FINAL=OUT/'林肯达_研究生学位论文中期考核表.docx'
doc=Document(REF); pp=list(doc.paragraphs); tables=list(doc.tables)
TITLE='面向大模型的数据与参数联合高效微调算法研究'

# Preserve the reference package and edit only its designated content slots.
def font(r,size=10.5,bold=False,ea='宋体'):
    r.font.name='Times New Roman'; r.font.size=Pt(size); r.bold=bold
    rp=r._r.get_or_add_rPr(); f=rp.find(qn('w:rFonts'))
    if f is None: f=OxmlElement('w:rFonts'); rp.insert(0,f)
    f.set(qn('w:eastAsia'),ea)
    r.font.color.rgb=None
    return r

def para(p,text='',size=10.5,bold=False,indent=True,align=WD_ALIGN_PARAGRAPH.JUSTIFY,before=0,after=3,line=16,keep=False,ea='宋体'):
    p.clear(); p.style=doc.styles['Normal']
    pr=p._p.get_or_add_pPr()
    # Eliminate source-specific character positioning and hidden numbering.
    for el in list(pr):
        if el.tag in [qn('w:numPr'),qn('w:tabs'),qn('w:ind'),qn('w:rPr'),qn('w:framePr')]: pr.remove(el)
    f=p.paragraph_format; f.alignment=align
    f.first_line_indent=Pt(21 if indent else 0); f.left_indent=Pt(0); f.right_indent=Pt(0)
    f.space_before=Pt(before); f.space_after=Pt(after)
    f.line_spacing=Pt(line); f.keep_with_next=keep; f.keep_together=False; f.widow_control=True
    f.page_break_before=False
    snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); pr.append(snap)
    font(p.add_run(text),size,bold,ea)
    return p

def blankcell(c):
    tc=c._tc
    for ch in list(tc):
        if ch.tag!=qn('w:tcPr'): tc.remove(ch)
    tc.append(OxmlElement('w:p'))
    return c.paragraphs[0]

def setcell(c,text='',**kwargs): return para(blankcell(c),text,**kwargs)
def addp(c,text='',**kwargs):
    p=c.paragraphs[0] if len(c.paragraphs)==1 and not c.paragraphs[0].text and len(c._tc)==2 else c.add_paragraph()
    return para(p,text,**kwargs)

def heading(c,text): return addp(c,text,bold=True,indent=False,before=7,after=4,keep=True,line=17)

def margins(table,top=85,start=110,bottom=85,end=110):
    pr=table._tbl.tblPr
    for tag in ['w:tblpPr','w:tblOverlap','w:tblInd']:
        for e in pr.findall(qn(tag)):pr.remove(e)
    m=pr.find(qn('w:tblCellMar'))
    if m is None: m=OxmlElement('w:tblCellMar'); pr.append(m)
    for side,val in [('top',top),('left',start),('bottom',bottom),('right',end)]:
        e=m.find(qn('w:'+side))
        if e is None: e=OxmlElement('w:'+side);m.append(e)
        e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa')
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    for row in table.rows:
        prr=row._tr.get_or_add_trPr()
        for e in list(prr):
            if e.tag in [qn('w:trHeight'),qn('w:cantSplit')]:prr.remove(e)
    for tc in table._tbl.iter(qn('w:tc')):
        prc=tc.find(qn('w:tcPr'))
        if prc is not None:
            for tag in ['w:tcMar','w:noWrap']:
                for e in prc.findall(qn(tag)): prc.remove(e)

def borders(t,color='000000',sz='4'):
    p=t._tbl.tblPr; el=p.find(qn('w:tblBorders'))
    if el is not None:p.remove(el)
    el=OxmlElement('w:tblBorders');p.append(el)
    for name in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+name); b.set(qn('w:val'),'single');b.set(qn('w:sz'),sz);b.set(qn('w:color'),color);el.append(b)

def smalltable(cell,headers,rows,widths):
    t=cell.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for c,w,label in zip(t.rows[0].cells,widths,headers):
        c.width=Cm(w);setcell(c,label,size=10,bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=15,after=0)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');c._tc.get_or_add_tcPr().append(sh)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for j,(c,w,value) in enumerate(zip(t.add_row().cells,widths,row)):
            c.width=Cm(w);setcell(c,value,size=10,indent=False,align=WD_ALIGN_PARAGRAPH.LEFT if len(value)>20 else WD_ALIGN_PARAGRAPH.CENTER,line=15,after=0)
    margins(t,65,90,65,90);borders(t,'808080')
    for row in t.rows:
        cs=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cs)
        for c in row.cells:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    # Word requires the paragraph that follows a nested table.
    para(cell.paragraphs[-1],size=1,indent=False,line=2,after=0)
    return t

# Cover keeps the original school logo bytes and visual order.
keep_cover={1,5,8,14,15,16,17,18,20,25,26}
for i in range(28):
    if i not in keep_cover:
        el=pp[i]._p;el.getparent().remove(el)
logo=pp[1]; logo.paragraph_format.space_before=Pt(15);logo.paragraph_format.space_after=Pt(40)
logo.paragraph_format.line_spacing=1;logo.paragraph_format.keep_with_next=True
logo.paragraph_format.first_line_indent=Pt(0)
if 'Title' not in doc.styles:doc.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
p=para(pp[5],'研究生学位（毕业）论文\n中期考核表',size=28,bold=True,ea='黑体',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=36,after=34,keep=True)
p.style=doc.styles['Title']
para(pp[8],'面向大模型的数据与参数\n联合高效微调算法研究',size=18,bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=28,after=40,keep=True)
fields=[('研 究 生','林肯达'),('指导教师','杨磊'),('学    号','202421045825'),('院（系）','软件学院'),('专    业','软件工程')]
identity=doc.add_table(rows=5,cols=2)
identity.alignment=WD_TABLE_ALIGNMENT.CENTER;identity.autofit=False
identity.columns[0].width=Cm(3.6);identity.columns[1].width=Cm(5.2)
# Use an invisible two-column grid for consistent field alignment.
pr=identity._tbl.tblPr
b=OxmlElement('w:tblBorders');pr.append(b)
for side in ['top','left','bottom','right','insideH','insideV']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');b.append(e)
for row,(label,value) in zip(identity.rows,fields):
    row.cells[0].width=Cm(3.6);row.cells[1].width=Cm(5.2)
    setcell(row.cells[0],label+'：',size=14,indent=False,align=WD_ALIGN_PARAGRAPH.RIGHT,line=22,after=6)
    setcell(row.cells[1],value,size=14,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=22,after=6)
    bottom=OxmlElement('w:pBdr');edge=OxmlElement('w:bottom');edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:space'),'1');bottom.append(edge)
    row.cells[1].paragraphs[0]._p.get_or_add_pPr().append(bottom)
    for c in row.cells:c.paragraphs[0].paragraph_format.keep_with_next=True
margins(identity,0,90,0,90)
pp[14]._p.addprevious(identity._tbl)
for idx in range(14,19):pp[idx]._p.getparent().remove(pp[idx]._p)
para(pp[20],'□ 博士研究生       ☑ 硕士研究生',size=14,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=21,after=47,line=23,keep=True)
para(pp[25],'华南理工大学研究生院',size=14,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=23,after=4,keep=True)
para(pp[26],'二〇二六年九月十四日',size=14,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=23,after=0)

# Major headings and sequential table flow.
for i,label in [(28,'一、中期考核工作基本情况'),(29,'二、中期考核作者承诺及导师意见'),(42,'三、中期考核审核意见')]:
    p=para(pp[i],label,size=12,bold=True,indent=False,line=20,after=9,keep=True,ea='仿宋_GB2312')
    p.paragraph_format.page_break_before=True
for i in range(30,42):
    el=pp[i]._p
    if el.getparent() is not None:el.getparent().remove(el)
for p in pp[43:]:
    if not p.text and not p._p.xpath('.//w:drawing'):
        el=p._p
        if el.getparent() is not None:el.getparent().remove(el)

main=tables[0]; margins(main);borders(main)
for row in main.rows:
    for c in row.cells:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for i,label in [(0,'开题报告题目'),(1,'文献综述题目')]:
    setcell(main.cell(i,0),label,indent=False,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,line=17,after=0)
    setcell(main.cell(i,1),TITLE if i==0 else '',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=17,after=0)
    main.rows[i].height=Cm(1.0);main.rows[i].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
for i,label in [(2,'课题研究的创新点：'),(4,'开题以来研究工作总结：'),(6,'与课题相关的学术成果情况：')]:
    p=setcell(main.cell(i,0),label,bold=True,indent=False,line=17,after=0,keep=True)

innovation=[
('1 背景与动机',[
'大语言模型通过指令微调适应下游任务，但大量候选数据会带来较高的训练与评测成本。参数高效微调（PEFT）通过仅更新少量参数降低训练资源需求，数据选择则通过保留更有价值的样本减少数据开销。两者共同影响训练过程：相同样本在不同目标任务、不同更新位置和参数容量下，可能具有不同的训练价值。因此，需要结合目标任务与 PEFT 配置进行数据选择。',
'现有数据选择方式主要面临两方面局限：语义相似度或统一样本评分难以充分表达配置差异；为每种配置构建专属梯度数据存储，并针对各任务分别评分，虽具有针对性，却会在新增配置时产生较高的重复开销。本课题以多配置场景为研究对象，探索可复用的条件效用估计与高效实验执行，使数据选择与参数更新方式相匹配。']),
('2 工作一已形成的方法设计',[
'（1）配置与任务共同条件化的样本效用建模。将样本、PEFT 配置与任务共同作为评分条件，借助共享干预位点建立可比较的信号空间，并以结构化编码区分更新位置、算子、容量及训练配方。在公共表征复用的基础上，为不同目标组合生成各自的候选排序。',
'（2）多保真监督驱动的可复用评分器。利用共享位点梯度代理提供较大范围的低成本监督，再以少量短更新观测校准代理排序；通过条件调制和交互建模学习样本价值及不确定性，降低为每种配置重新构建完整选择器的需求。',
'（3）效用与语义覆盖兼顾的预算分配。根据各目标的保守评分估计语义簇价值，结合簇规模分配选数配额，缓解全局高分样本过于集中的问题。通过区分离线建模与新增配置开销，研究目标针对性与多配置成本摊销之间的关系。']),
('3 工作二拟开展的工程机制研究',[
'在 PCU-Select 方法与实现基础上，拟研究分层缓存与增量复用、公共计算去重与阶段恢复两项机制，构建多配置数据选择与微调实验管理系统。重点解决输入变化后的正确失效、公共产物复用、批量实验隔离和中断恢复问题，以正确性、重复计算量、运行时间和恢复开销进行定量验证。该部分属于下一阶段工作计划。'])]
ic=main.cell(3,0);blankcell(ic)
for h,ps in innovation:
    heading(ic,h)
    for text in ps:addp(ic,text)

work=json.loads((TMP/'work1_content.json').read_text())
for sec in work:
    sec['paragraphs']=[p.replace('96 维描述 24 个位点上的激活状态和更新算子','96 维描述 24 个位点上的有效位点掩码和更新算子').replace('不同配置、任务、起点和步数分别归一化','在同一配置、任务、起点与更新步数构成的条件组内进行排序归一化').replace('价值与规模的指数分别为 0.6 和 0.4','先对簇价值取正部，再按价值与规模分别取 0.6 和 0.4 次幂').replace('设 P 为需要服务的 PEFT 配置数','对无须额外校准的受支持配置，设 P 为需要服务的 PEFT 配置数') for p in sec['paragraphs']]
work[5]['paragraphs'][0]=work[5]['paragraphs'][0].replace('现有汇总不能代替原始运行证据，也不作为已完成多次独立复现实验的证明。','结合独立重复结果与任务差异，说明结论的适用范围。')
wc=main.cell(5,0);blankcell(wc)
for sec in work:
    heading(wc,sec['heading'])
    for pi,text in enumerate(sec['paragraphs']):
        addp(wc,text)
        if sec['heading'].startswith('5 ') and pi==1:
            addp(wc,'表 1  当前稿件主表宏平均汇总（四任务 × 五配置）',bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=3,after=4,keep=True)
            smalltable(wc,['方法','宏平均','相对 PCU-Select 的差值'],[
                ['Random','32.29','−2.89'],['RDS+','34.44','−0.74'],['Influence','34.68','−0.50'],['逐配置 LESS','35.14','−0.04'],['PCU-Select','35.18','—']],[5.7,3.8,5.7])
            addp(wc,'注：不同任务原生百分制指标的宏平均，仅用于当前统一实验协议下的汇总比较。',size=9,indent=False,line=14,after=4)

system=[
('7 工作二：系统研究内容与实施方案',[
'第二项工作拟开展“基于 PCU-Select 的多配置数据选择与微调实验管理系统”研究，直接复用第一项工作的评分器、样本特征与任务表征、PEFT 编码、聚类配额算法及训练评测脚本。研究重点从单次选择扩展到反复开展多任务、多预算、多训练种子实验时的产物管理与执行效率，预期形成一套可运行系统和两项可定量验证的工程机制。',
'（1）分层缓存与增量复用。为样本特征、任务表示、评分、聚类及选择结果建立依赖清单，记录样本内容摘要、模型与评分器版本、预处理配置、随机种子、候选顺序和产物校验值。输入变化后沿依赖关系只重算受影响节点，避免同名样本内容变化或预算不同导致错误命中旧结果。首版固定评分器版本，并将版本显式写入产物依赖。',
'复用规则包括：仅改变选数预算时，复用评分和聚类，重新计算配额与子集；选择种子不变而仅改变下游训练种子时，复用选择结果。增加候选样本且评分条件不变时，复用旧样本的兼容特征与逐样本评分，仅补算新增样本；完整候选池的聚类、配额和子集仍需重算。候选顺序变化亦触发相应依赖检查，以可追溯的输入标识约束复用范围。',
'（2）公共计算去重与阶段恢复。将实验请求拆解为特征、任务表示、评分、聚类、选数、训练和评测节点，输入依赖完全相同的节点仅执行一次。采用临时文件原子写入、校验与完成标记，结合请求独立目录，避免半成品被复用及结果相互覆盖。中断后保留已完成且校验通过的产物，仅重跑失败或未完成阶段；首版实现阶段级恢复。',
'以同一数据池上的两种 PEFT、两个任务、三个预算及两个训练种子为例，共形成 24 次训练实验。固定特征与评分器版本、聚类配置及选择种子后，计划共享四份目标评分和一份聚类，生成十二份目标子集，再分别完成二十四次训练。各任务和配置保留自己的子集；公共节点数量减少并不直接等同于端到端加速倍数。',
'首版范围限定为单工作节点、任务队列、元数据索引和本地产物库，并提供简洁控制台。操作流程包括登记数据与任务、提交实验矩阵、查看复用计划、执行选数、按需训练评测及导出报告。继续沿用第一项工作已验证的骨干与 PEFT 支持范围，优先完成上述机制的正确性与有效性验证。']),
('8 验证方案与下一阶段工作计划',[
'系统验证拟设置现有批量脚本、增加依赖缓存、完整系统三组方案，公平保留现有脚本已具备的特征缓存能力。对比冷启动、多预算和多训练种子复用、追加 1%、5%、10% 候选数据等场景，并检查同名样本内容变化、候选顺序变化等失效情况。复用所得评分和子集需与相同输入下的全量重算结果一致，再统计节点执行次数、选择阶段时间、端到端时间、存储占用及校验管理开销。',
'恢复验证拟在评分、子集写入和训练等阶段注入中断，检查未完成产物是否被隔离、已完成公共节点是否正确复用、失败请求是否影响其他实验，以及恢复所需的重算量。工程收益以实测结果报告，不预设固定加速倍数。工作一同步推进原始日志核验、独立重复实验、任务差异分析及需要额外校准场景的成本补充。',
'下一阶段拟以 2026 年 9 月 14 日至 11 月 8 日为八周实施周期，按下表推进。各阶段以可运行产物和可核验记录为验收依据，并根据实验资源与导师指导适当调整。'])]
for h,ps in system:
    heading(wc,h)
    for text in ps:addp(wc,text)
addp(wc,'表 2  下一阶段工作安排',bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=3,after=4,keep=True)
planrows=[
['第 1—2 周\n9.14—9.27','梳理现有接口与产物；设计依赖键、目录和元数据；整理工作一实验记录。','完成依赖规范及最小执行流程，建立稿件结果核验清单。'],
['第 3—4 周\n9.28—10.11','实现分层缓存与增量计算；验证预算、种子、样本追加和内容变化规则。','完成缓存原型及失效验证，形成全量重算一致性记录。'],
['第 5—6 周\n10.12—10.25','实现公共节点去重、任务队列、阶段恢复及控制台；运行 24 次训练实验矩阵。','形成可运行系统，完成隔离与中断恢复测试。'],
['第 7—8 周\n10.26—11.8','完成三组对比与开销分析；整理工作一实验记录；整理系统报告与学位论文。','形成机制验证结果、系统说明和学位论文相关章节。']]
smalltable(wc,['时间','主要工作','阶段产出'],planrows,[2.5,6.4,6.3])
addp(wc,'预期在上述阶段形成完整的算法—系统研究链条：以工作一回答不同配置应选择哪些数据，以工作二保证多配置实验中公共计算可复用、结果可追溯、失败可恢复，并以统一实验记录支撑学位论文撰写。',before=4)

ac=main.cell(7,0);blankcell(ac)
addp(ac,'已形成英文论文稿《PCU-Select: Amortizing Target-Specific Data Selection Across PEFT Configurations》，围绕条件效用建模、共享干预位点、多保真监督和多配置成本摊销组织了方法与实验章节，目前继续完善原始实验记录核验、结果分析与论证。')
addp(ac,'已形成与论文配套的算法实现和训练评测脚本基础，涵盖特征提取、PEFT 配置注册、条件评分、聚类配额选数及结果汇总，为第二项系统工作提供可复用模块。第二项系统及其机制验证属于计划成果，尚待下一阶段完成。')

# Keep each field label inside its long content cell, so Word can flow
# the body without treating the next multi-page table row as one keep group.
original_rows=list(main.rows)
for label_index,body_index in [(2,3),(4,5),(6,7)]:
    label_node=original_rows[label_index].cells[0].paragraphs[0]._p
    body_cell=original_rows[body_index].cells[0]
    body_cell._tc.insert(1,deepcopy(label_node))
    body_cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    cp=body_cell._tc.get_or_add_tcPr()
    for e in cp.findall(qn('w:tcBorders')):cp.remove(e)
    main._tbl.remove(original_rows[label_index]._tr)

for row in main.rows:
    for e in list(row._tr):
        if e.tag==qn('w:tblPrEx'):row._tr.remove(e)
    cs=OxmlElement('w:cantSplit');cs.set(qn('w:val'),'0');row._tr.get_or_add_trPr().append(cs)
    for c in row.cells:
        for p in c.paragraphs[:2]:p.paragraph_format.keep_with_next=False

for p in wc.paragraphs:
    if p.text=='8 验证方案与下一阶段工作计划':p.paragraph_format.space_before=Pt(30)

# Printed declarations are retained; actual signatures and opinion fields stay empty.
commit=tables[1];margins(commit,150,140,130,140);borders(commit)
texts=['我保证上述填报内容的真实性，并将在导师指导下，严格遵守学校的有关规定，按计划认真开展学位（毕业）论文研究工作。','我已审阅过中期考核的全部内容，同意研究生参加中期考核。']
for i in range(2):
    c=commit.cell(i,0);blankcell(c);addp(c,texts[i],size=12,line=23,after=28)
    addp(c,('研究生签名：' if i==0 else '指导教师签名：')+'________________',size=12,indent=False,align=WD_ALIGN_PARAGRAPH.RIGHT,line=23,after=15)
    addp(c,'年    月    日',size=12,indent=False,align=WD_ALIGN_PARAGRAPH.RIGHT,line=23,after=0)
    commit.rows[i].height=Cm(5.3);commit.rows[i].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    cs=OxmlElement('w:cantSplit');commit.rows[i]._tr.get_or_add_trPr().append(cs)

review=tables[2];margins(review,50,90,50,90);borders(review)
# Normalize every physical cell while preserving all existing grid spans and vertical merges.
for tc in review._tbl.iterchildren(qn('w:tr')):
    for cnode in tc.iterchildren(qn('w:tc')):
        from docx.table import _Cell
        c=_Cell(cnode,review)
        values=[p.text.strip() for p in c.paragraphs if p.text.strip()]
        # Preserve only printed field labels; no reference author's signature/opinion exists.
        p=blankcell(c)
        for k,text in enumerate(values):
            para(p if k==0 else c.add_paragraph(),text,size=10.5,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER if text not in ['具体意见：','审核结果：','组长签名：','成员签名：'] else WD_ALIGN_PARAGRAPH.LEFT,line=17,after=0)
        if not values:para(p,'',indent=False,line=12,after=0)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
# Rows 10 and 17 are blank writing areas; use source labels and merged geometry as-is.
heights=[.75,.75]+[.72]*8+[.18,3.7,.7,.75,1.0,1.05,.65,2.25,.8,.65]
for row,h in zip(review.rows,heights):
    row.height=Cm(h);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    cs=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cs)
# Put specific-opinion prompt at the top of its writing space.
for tc in review._tbl.iter(qn('w:tc')):
    from docx.table import _Cell
    c=_Cell(tc,review)
    if c.text.startswith('具体意见：'):
        setcell(c,'具体意见：',size=10.5,indent=False,align=WD_ALIGN_PARAGRAPH.LEFT,line=18,after=0)
        addp(c,'审核结果：',size=10.5,indent=False,align=WD_ALIGN_PARAGRAPH.LEFT,before=74,line=18,after=0)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP

sec=doc.sections[0];sec.different_first_page_header_footer=True
num=sec._sectPr.find(qn('w:pgNumType'));num.set(qn('w:fmt'),'decimal')
foot=sec.footer.paragraphs[0];para(foot,'',size=10.5,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=15,after=0)
r=font(foot.add_run(),10.5)
fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),'begin');r._r.append(fld)
r=font(foot.add_run(),10.5);instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' PAGE ';r._r.append(instr)
r=font(foot.add_run(),10.5);sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
font(foot.add_run('1'),10.5)
r=font(foot.add_run(),10.5);end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)
settings=doc.settings.element
for tag in ['w:docVars','w:rsids']:
    for e in settings.findall(qn(tag)):settings.remove(e)
up=settings.find(qn('w:updateFields'))
if up is None:up=OxmlElement('w:updateFields');settings.append(up)
up.set(qn('w:val'),'true')
doc.core_properties.author='林肯达';doc.core_properties.last_modified_by='林肯达';doc.core_properties.title='研究生学位（毕业）论文中期考核表';doc.core_properties.subject=TITLE
doc.core_properties.comments='';doc.core_properties.keywords='PCU-Select；中期考核'
intermediate=TMP/'authored.docx';doc.save(intermediate)

# Surgical ZIP assembly: preserve non-edited reference parts byte for byte.
allow={'word/document.xml','word/styles.xml','word/settings.xml','word/footer1.xml','docProps/core.xml','docProps/app.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with ZipFile(REF) as old,ZipFile(intermediate) as new:
    original={n:old.read(n) for n in old.namelist()}; revised={n:new.read(n) for n in new.namelist()}
    removed={n for n in original if n.startswith('word/media/') and n!='word/media/image1.jpeg'}
    rel=etree.fromstring(original['word/_rels/document.xml.rels'])
    for r in list(rel):
        if 'word/'+r.get('Target','') in removed:rel.remove(r)
    revised['word/_rels/document.xml.rels']=etree.tostring(rel,encoding='UTF-8',xml_declaration=True,standalone=True)
    ct=etree.fromstring(original['[Content_Types].xml'])
    for e in list(ct):
        if e.get('PartName','').lstrip('/') in removed:ct.remove(e)
    revised['[Content_Types].xml']=etree.tostring(ct,encoding='UTF-8',xml_declaration=True,standalone=True)
    app=etree.fromstring(original['docProps/app.xml'])
    for el in app.iter():
        if etree.QName(el).localname in ['Company','Manager']:el.text=''
    revised['docProps/app.xml']=etree.tostring(app,encoding='UTF-8',xml_declaration=True,standalone=True)
    with ZipFile(FINAL,'w',ZIP_DEFLATED) as z:
        for n,b in original.items():
            if n not in removed:z.writestr(n,revised.get(n,b) if n in allow else b)
    manifest={'source':str(REF),'source_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'output':str(FINAL),'edited_parts':[],'preserved_parts':[],'removed_parts':sorted(removed)}
    with ZipFile(FINAL) as z:
        for n in z.namelist():
            (manifest['preserved_parts'] if z.read(n)==original[n] else manifest['edited_parts']).append(n)
            if n not in allow:assert z.read(n)==original[n],n
        assert z.read('word/media/image1.jpeg')==original['word/media/image1.jpeg']
        txt=z.read('word/document.xml').decode()
        for value in ['林肯达','202421045825','杨磊','软件学院','软件工程','二〇二六年九月十四日']:assert value in txt,value
        for value in ['陈瑾如','陈浩锐','卢彦谚','202121047085','移动边缘网络']:assert value not in txt,value
    (TMP/'package_validation.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
(TMP/'final_content.json').write_text(json.dumps({'innovation':innovation,'work1':work,'system':system,'schedule':planrows},ensure_ascii=False,indent=2))
print(FINAL)
print('Output bytes:',FINAL.stat().st_size)
print('Edited parts:',manifest['edited_parts'])
