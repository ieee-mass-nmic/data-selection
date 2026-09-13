from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree
import json,ast,hashlib
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path('/Users/bytedance/code/data-selection');TMP=ROOT/'tmp/midterm_assessment_revision'
BASE=ROOT/'output/midterm_assessment/林肯达_中期考核表_论文式图文版.docx'
OUT=ROOT/'output/midterm_assessment/林肯达_中期考核表_公式与正文实验修订版.docx'
doc=Document(BASE)
# Reuse the existing reference-derived typography and table helpers.
module=ast.parse((ROOT/'tmp/midterm_assessment/build_form.py').read_text())
for n in module.body:
 if isinstance(n,ast.FunctionDef):exec(compile(ast.Module(body=[n],type_ignores=[]),'reference_helpers','exec'))
main=doc.tables[1];ic=main.cell(2,0);wc=main.cell(3,0);ac=main.cell(4,0)
method=json.loads((ROOT/'tmp/midterm_assessment_paper/method_content.json').read_text());sec={s['key']:s for s in method['sections']}
review=json.loads((ROOT/'tmp/midterm_assessment_paper/figure_review.json').read_text())

import re

def E(tag,*children,**attrs):
    el=OxmlElement('m:'+tag)
    for k,v in attrs.items():el.set(qn('m:'+k),str(v))
    for c in children:
        if isinstance(c,list):
            for cc in c:el.append(cc)
        else:el.append(c)
    return el

def R(text,normal=False):
    pr=E('rPr')
    if normal and any(c.isascii() and c.isalpha() for c in text):pr.append(E('nor'))
    pr.append(E('sty',val='p' if normal else 'i'))
    r=E('r',pr)
    wp=OxmlElement('w:rPr')
    f=OxmlElement('w:rFonts')
    for name in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+name),'STIX Two Math')
    wp.append(f)
    for name in ['i','iCs']:
        item=OxmlElement('w:'+name);item.set(qn('w:val'),'0' if normal else '1');wp.append(item)
    for name in ['sz','szCs']:
        item=OxmlElement('w:'+name);item.set(qn('w:val'),'22');wp.append(item)
    r.append(wp)
    t=E('t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t)
    return r

def V(s):return R(s)
def N(s):
    if s in ['=','+','−','≈']:s='\u2009'+s+'\u2009'
    if s=='−0.2':s='\u2009−\u20090.2'
    return R(s,True)
def seq(x):return x if isinstance(x,list) else [x if not isinstance(x,str) else V(x)]
def sub(b,s):return E('sSub',E('e',*seq(b)),E('sub',*seq(s)))
def sup(b,s):return E('sSup',E('e',*seq(b)),E('sup',*seq(s)))
def ss(b,s,u):return E('sSubSup',E('e',*seq(b)),E('sub',*seq(s)),E('sup',*seq(u)))
def frac(n,d):return E('f',E('num',*seq(n)),E('den',*seq(d)))
def group(xs,left='(',right=')'):
    return E('d',E('dPr',E('begChr',val=left),E('endChr',val=right)),E('e',*seq(xs)))
def nsum(low,term):
    return E('nary',E('naryPr',E('chr',val='∑'),E('limLoc',val='subSup'),E('supHide',val='1')),E('sub',*seq(low)),E('sup'),E('e',*seq(term)))
def hat(x):return E('acc',E('accPr',E('chr',val='̂')),E('e',*seq(x)))
def tilde(x):return E('acc',E('accPr',E('chr',val='̃')),E('e',*seq(x)))
def fun(name,args):return E('func',E('fName',N(name)),E('e',group(args)))
def comma():return N(',')

def eqnodes(key):
    if key=='selection_objective':
        constraint=[V('S'),N('⊆'),V('D'),N(', '),group('S','|','|'),N('≤'),V('B')]
        arg=E('limLow',E('e',N('arg max')),E('lim',*constraint))
        value=fun('Perf',[fun('Tune',[sub('M',N('0')),comma(),V('p'),comma(),V('S')]),comma(),ss('V','t',N('test'))])
        return [sup('S',N('*')),N('='),E('func',E('fName',arg),E('e',value))]
    if key=='proxy_label':
        return [sup('u',N('lo')),group([V('x'),comma(),V('p'),comma(),V('t')]),N('='),nsum([V('ω'),N('∈'),V('Ω')],[ss(tilde('w'),'p','ω'),fun('cos',[ss('g','x','ω'),comma(),ss('g','t','ω')])])]
    if key=='short_update':
        theta=sub('θ','a');lv=sub('L',sub('V','t'))
        adapt=E('func',E('fName',ss(N('Adapt'),'p','h')),E('e',group([V('B'),group('x'),N(';'),deepcopy(theta)])))
        return [V('Δ'),group([V('x'),comma(),V('p'),comma(),V('t'),comma(),V('a'),comma(),V('h')]),N('='),deepcopy(lv),group(deepcopy(theta)),N('−'),deepcopy(lv),group(adapt)]
    if key=='conservative_score':
        args=group([V('x'),comma(),sup('p',N('*')),comma(),sup('t',N('*'))])
        return [V('q'),group('x'),N('='),sub(hat('μ'),'φ'),deepcopy(args),N('−0.2'),sub(hat('σ'),'φ'),deepcopy(args)]
    if key=='cluster_quota':
        def value(i):return [sup(group(ss('v',i,N('+'))),N('0.6')),sup(group(sub('C',i),'|','|'),N('0.4'))]
        return [sub('b','k'),N('='),E('func',E('fName',N('round')),E('e',group([V('B'),frac(value('k'),nsum('j',value('j')))],'[',']'))),N(',   '),ss('v','k',N('+')),N('='),fun('max',[sub('v','k'),comma(),N('0')])]
    if key=='cost_general':
        lines=[
            [sub('C',N('PCU')),group([V('P'),comma(),V('Q')]),N('='),sub('C',N('offline')),N('+'),V('PQ'),ss('C',N('PCU'),N('pair'))],
            [sub('C',N('LESS')),group([V('P'),comma(),V('Q')]),N('='),V('P'),ss('C',N('LESS'),N('data')),N('+'),V('PQ'),ss('C',N('LESS'),N('rank'))]]
        return [E('eqArr',E('eqArrPr',E('baseJc',val='center')), *[E('e',*line) for line in lines])]
    if key=='break_even':return [sup('P',N('*')),N('='),frac(N('144'),N('63.2−1.51')),N('≈2.33')]
    raise KeyError(key)

def hidden_borders(t):
    for b in list(t._tbl.tblPr.findall(qn('w:tblBorders'))):t._tbl.tblPr.remove(b)
    b=OxmlElement('w:tblBorders');t._tbl.tblPr.append(b)
    for name in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+name);edge.set(qn('w:val'),'nil');b.append(edge)

eqcount=0
def equation(key):
    global eqcount
    eqcount+=1
    t=wc.add_table(rows=1,cols=3);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,c,w in zip(t.columns,t.rows[0].cells,[0.75,13.5,0.75]):
        col.width=Cm(w);c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(t,60,0,60,0);hidden_borders(t)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for c in t.rows[0].cells:
        p=setcell(c,'',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,after=0)
        p.paragraph_format.line_spacing=1.15;p.paragraph_format.keep_together=True
    p=t.cell(0,1).paragraphs[0]
    p._p.append(E('oMathPara',E('oMathParaPr',E('jc',val='center')),E('oMath',*eqnodes(key))))
    p=t.cell(0,2).paragraphs[0];p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(p.add_run('('+str(eqcount)+')'),10.5)
    para(wc.paragraphs[-1],'',size=1,line=2,indent=False,after=2)
    return t

def heading2(c,text):return addp(c,text,bold=True,indent=False,before=7,after=4,line=17,keep=True)
def major(c,text):return addp(c,text,bold=True,indent=False,before=9,after=5,line=18,keep=True)

def fig(path,caption,width,cell=None):
    target=wc if cell is None else cell
    t=target.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    t.columns[0].width=Cm(15.0);t.cell(0,0).width=Cm(15.0)
    margins(t,30,0,30,0);hidden_borders(t)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    c=t.cell(0,0);p=setcell(c,'',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,after=3)
    p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
    pic=p.add_run().add_picture(str(path),width=Cm(width));pic._inline.docPr.set('descr',caption)
    para(c.add_paragraph(),caption,size=9.5,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=14,after=0)
    para(target.paragraphs[-1],'',size=1,line=2,indent=False,after=0)
    return t

PPR_ORDER='pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange'.split()
def normalize_body_math_and_properties():
    mp=doc.settings.element.find(qn('m:mathPr'))
    if mp is None:mp=E('mathPr');doc.settings.element.append(mp)
    mf=mp.find(qn('m:mathFont'))
    if mf is None:mf=E('mathFont');mp.insert(0,mf)
    mf.set(qn('m:val'),'STIX Two Math')
    names={s:i for i,s in enumerate(PPR_ORDER)}
    # Only the newly authored main form changes; cover and signatures retain exact XML.
    for pr in main._tbl.iter(qn('w:pPr')):
        children=list(pr)
        for e in children:pr.remove(e)
        for e in sorted(children,key=lambda e:names.get(etree.QName(e).localname,100)):pr.append(e)
    table_order='tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange'.split()
    order={s:i for i,s in enumerate(table_order)}
    for pr in main._tbl.iter(qn('w:tblPr')):
        if pr.getparent() is main._tbl:continue
        children=list(pr)
        for e in children:pr.remove(e)
        for e in sorted(children,key=lambda e:order.get(etree.QName(e).localname,100)):pr.append(e)
    # Keep the achievements label and its response on the same page.
    tr=main.rows[4]._tr.get_or_add_trPr()
    for item in list(tr.findall(qn('w:cantSplit'))):tr.remove(item)
    tr.append(OxmlElement('w:cantSplit'))
    ac.paragraphs[0].paragraph_format.keep_with_next=True
    symbols={'M₀':lambda:sub('M',N('0')),'Vₜ':lambda:sub('V','t'),'θₐ':lambda:sub('θ','a'),'μ̂':lambda:hat('μ'),'σ̂':lambda:hat('σ'),'Cₖ':lambda:sub('C','k'),'vₖ⁺':lambda:ss('v','k',N('+')),'vₖ':lambda:sub('v','k')}
    pat=re.compile('('+'|'.join(re.escape(s) for s in symbols)+')')
    for run in list(main._tbl.iter(qn('w:r'))):
        txt=''.join(run.xpath('./w:t/text()'))
        if not pat.search(txt):continue
        parent=run.getparent();where=parent.index(run);rp=run.find(qn('w:rPr'))
        for part in pat.split(txt):
            if not part:continue
            if part in symbols:node=E('oMath',symbols[part]())
            else:
                node=OxmlElement('w:r')
                if rp is not None:node.append(deepcopy(rp))
                text=OxmlElement('w:t');text.text=part;text.set(qn('xml:space'),'preserve');node.append(text)
            parent.insert(where,node);where+=1
        parent.remove(run)


def keep_caption_with_table(table):
    tbl = table._tbl
    caption = tbl.getprevious()
    parent = tbl.getparent()
    width = sum(int(col.get(qn('w:w'))) for col in tbl.find(qn('w:tblGrid')))
    wrapper = OxmlElement('w:tbl')
    props = OxmlElement('w:tblPr')
    wrapper.append(props)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(width)); w.set(qn('w:type'), 'dxa'); props.append(w)
    align = OxmlElement('w:jc'); align.set(qn('w:val'), 'center'); props.append(align)
    borders = OxmlElement('w:tblBorders'); props.append(borders)
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + side); border.set(qn('w:val'), 'nil'); borders.append(border)
    margins = OxmlElement('w:tblCellMar'); props.append(margins)
    for side in ('top', 'left', 'bottom', 'right'):
        margin = OxmlElement('w:' + side); margin.set(qn('w:w'), '0'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
    grid = OxmlElement('w:tblGrid'); col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(width)); grid.append(col); wrapper.append(grid)
    row = OxmlElement('w:tr'); row_props = OxmlElement('w:trPr'); row_props.append(OxmlElement('w:cantSplit')); row.append(row_props); wrapper.append(row)
    cell = OxmlElement('w:tc'); cell_props = OxmlElement('w:tcPr'); cell_width = OxmlElement('w:tcW'); cell_width.set(qn('w:w'), str(width)); cell_width.set(qn('w:type'), 'dxa'); cell_props.append(cell_width); cell.append(cell_props); row.append(cell)
    for page_break in list(caption.iter(qn('w:pageBreakBefore'))):
        page_break.set(qn('w:val'), '0')
    parent.insert(parent.index(caption), wrapper)
    cell.append(caption); cell.append(tbl)
    end = OxmlElement('w:p'); end_props = OxmlElement('w:pPr'); spacing = OxmlElement('w:spacing'); spacing.set(qn('w:line'), '20'); spacing.set(qn('w:lineRule'), 'exact'); end_props.append(spacing); end.append(end_props); cell.append(end)

# The official form labels remain; the content now follows a paper's logic.
blankcell(ic)
addp(ic,'课题研究的创新点：',bold=True,indent=False,line=17,after=4)
major(ic,'1 背景及动机')
for s in [
'大语言模型的指令微调能够提升模型对下游任务的适应能力，但大规模候选数据和反复试验会带来较高的训练成本。参数高效微调（PEFT）通过限制可训练参数降低资源消耗，数据选择通过保留更有价值的样本减少数据开销。两类方法共同作用于训练过程：不同 PEFT 配置的更新位置、算子及容量不同，同一样本对目标任务的训练价值也可能随之变化。因而，数据选择需要同时考虑“训练什么任务”和“以何种方式更新模型”。',
'现有方法难以同时兼顾目标针对性和多配置复用。基于语义相似度或共享梯度的选择方式可以复用公共计算，但未充分表达具体 PEFT 配置的差异。逐配置 LESS 为每种配置构建专属梯度数据存储，再针对各任务评分，具有较强的目标针对性，却需要在新增配置时重复支付建库成本。直接沿用已有配置选出的子集，又可能损失与新配置的匹配性。本文由此研究可复用的条件效用评分：复用样本信号和评分器，并为不同配置与任务生成各自的训练子集。']:
 addp(ic,s)

heading2(ic,'1.1 背景动机实验')
addp(ic,'动机实验考察同一候选样本在不同 PEFT 配置下是否具有相同效用排序。实验设计固定候选池和目标任务，分别构造各配置的短更新效用排序，以 Spearman 相关系数比较整体排序，以前 5% 样本集合的 Jaccard 重叠衡量高价值样本的一致性；同配置的独立重复用于提供噪声参照。进一步将源配置选出的子集用于其他目标配置，检查排序差异是否影响下游训练效果。')
fig(TMP/'assets/motivation.png','图 1  不同 PEFT 配置的效用排序与前 5% 样本重叠',14.2,cell=ic)
addp(ic,'注：图源为论文引言中的动机分析图；生成脚本对非对角元素统一加 0.10。',size=9,indent=False,line=14,after=4)
addp(ic,'图 1 显示的趋势是：配置结构越接近，排序与高价值样本越一致；跨家族变化则可能产生更大差异。该结果支持的研究假设是：数据价值应同时依赖目标任务和可训练子空间，复用公共计算时仍应保留配置条件。')

major(ic,'2 创新点')
for s in [
'（1）面向任务与参数配置的条件效用建模。将样本、PEFT 配置与任务共同作为效用预测条件，利用共享干预位点建立可比较的信号空间，并以结构化编码区分更新位置、算子、容量和训练配方，使公共信号复用与配置差异表达相结合。',
'（2）多保真监督驱动的统一评分器。以可缓存的位点梯度代理提供低成本排序监督，再通过少量短更新观测校正代理偏差；借助条件调制和交互建模学习样本效用与不确定性，减少重复构建目标专属选择器的成本。',
'（3）兼顾效用与覆盖的目标子集构建。将保守条件评分与语义簇规模共同用于预算分配，缓解全局高分样本过于集中的问题；通过区分一次性离线成本和新增配置成本，分析多配置数据选择的摊销收益。']:
 addp(ic,s)

blankcell(wc);addp(wc,'开题以来研究工作总结：',bold=True,indent=False,line=17,after=4)

def blocks(section):
 for b in sec[section]['blocks']:
  if b['type']=='paragraph':addp(wc,b['text'])
  elif b['type']=='equation':equation(b['id'])
  elif b['type']=='figure_reference':
   fig(ROOT/b['asset'],'图 2  PCU-Select 总体架构：共享信号提取、多保真效用学习与目标条件化选择',14.2)

major(wc,'1 问题定义');blocks('problem')
major(wc,'2 方法')
heading2(wc,'2.1 总体框架');blocks('framework')
heading2(wc,'2.2 共享干预位点与结构化表征');blocks('representations')
heading2(wc,'2.3 多保真效用监督');blocks('supervision')
heading2(wc,'2.4 条件评分与聚类配额选择');blocks('selection')
addp(wc,'评分器在已见配置及结构接近的目标上直接使用；对同一家族内较大偏移或未见家族，存在兼容标签时进行小规模校准，否则采用逐目标选择方法。该支持策略将可复用评分与适用边界结合，避免将固定骨干上的共享表征直接解释为任意配置、任意骨干均可迁移。')


major(wc,'3 实验')
heading2(wc,'3.1 实验设置')
for text in [
'实验按论文正文的三个问题展开：已见配置上的下游质量是否接近逐配置 LESS，公共计算能否降低新增配置的选择成本，以及评分器在未见配置上如何通过少量标签校准。采用 Llama-2-7B 骨干，从去污染后的 30 万条混合指令数据中选择 10%，即 3 万条样本。主实验覆盖 GSM8K、HumanEval、MMLU 和 TyDiQA，分别使用精确匹配率、Pass@1、准确率和 F1，并对四项百分制指标等权平均。',
'每个任务摘要含 32 个训练或开发侧示例，且与评测数据分离；HumanEval 使用 MBPP 和 CodeAlpaca 构造摘要。主实验采用 AD-b64、IA3-attnmlp、L-r16-qkvo、L-r8-mlp、L-r8-qv 五种已见配置。各选择方法固定骨干、PEFT 配方及评测协议，下游训练使用 1000 步 AdamW、全局批量 128、最大序列长度 1024 和余弦学习率调度。计时平台为 8 张 NVIDIA H20-96GB，GPU 小时等于墙钟小时乘以 GPU 数。',
'以下按照论文正文整理实验结果。表中转录均值；描述性统计不用于独立样本显著性结论。']:
 addp(wc,text)

heading2(wc,'3.2 实验一 已见配置上的下游质量对比')
addp(wc,'本实验比较条件选数与逐配置选择的下游质量。Random 不使用任务与 PEFT 条件；RDS+ 按冻结表征与任务摘要的相似度选数；Influence 使用不区分 PEFT 的共享梯度存储；LESS 为每种配置单独建库。PCU-Select 复用表征和评分器，为各目标生成子集。')
addp(wc,'表 1  五种已见 PEFT 配置的主实验结果',bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=4,after=4,keep=True)
edata=json.loads((ROOT/'tmp/midterm_assessment_paper/experiment_revision_review.json').read_text())
qrows=[[row[0]]+[f'{x:.2f}' for x in row[1:]] for row in edata['experiment_1_quality']['means']]
t=smalltable(wc,['方法','AD-b64','IA3-\nattnmlp','L-r16-\nqkvo','L-r8-\nmlp','L-r8-qv','平均'],qrows,[2.6,1.9,2.2,2.2,1.9,1.9,2.1])
keep_caption_with_table(t)
for c in t.rows[-1].cells:
 for p in c.paragraphs:
  for r in p.runs:r.bold=True
addp(wc,'注：改编自论文正文主实验表，仅保留均值。每个配置列为四个任务原生百分制指标的等权平均，最后一列对五种配置取平均。',size=9,indent=False,line=14,after=4)
addp(wc,'表 1 中 PCU-Select 的总体平均为 35.18，LESS 为 35.14，两者相差 0.04 个指标点；相对 RDS+ 和 Influence 分别增加 0.74 和 0.50 个点。该结果支持在当前注册集合上“平均质量接近”的阶段性判断。配置间仍存在差异：例如在 L-r8-mlp 上，PCU-Select 为 34.56，低于 LESS 的 35.13，因此不能将宏平均接近解释为所有配置和任务均占优。')

heading2(wc,'3.3 实验二 跨配置复用与摊销成本')
addp(wc,'本实验区分只支付一次的固定或共享成本，以及每新增一种 PEFT 配置的边际成本；每种配置均服务四个任务。除前述方法外，加入 Reuse-one-LESS：把一个来源配置选出的任务子集复用于所有目标配置，用于检验直接复用子集的质量代价。表 2 对照质量和选择成本，均排除共同的下游微调开销。')
addp(wc,'表 2  五种已见配置的质量与选择成本',bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=4,after=4,keep=True)
crows=[]
for row in edata['experiment_2_cost']['rows']:
 name,fixed,marg,total,quality,gap=row
 crows.append([name,f'{fixed:.1f}',f'{marg:.2f}',f'{total:.1f}',f'{quality:.2f}',f'{gap:+.2f}' if gap else '0.00'])
t=smalltable(wc,['方法','固定/共享\n成本','新增配置\n成本','五配置\n总成本','质量','相对 LESS\n差值'],crows,[3.8,2.2,2.3,2.3,2.0,2.2])
for c in t.rows[-1].cells:
 for p in c.paragraphs:
  for r in p.runs:r.bold=True
addp(wc,'注：改编自论文正文质量与成本表，成本单位为 GPU 小时；质量为四任务与五配置的宏平均，差值单位为指标点。该表按既定五配置、四任务集合计账。',size=9,indent=False,line=14,after=4)
addp(wc,'PCU-Select 将主要成本集中到共享离线阶段：一次性开销为 144.0 GPU 小时，新增四任务配置的成本为 1.51 GPU 小时；LESS 每个配置为 63.2 GPU 小时，边际选择成本比约为 42。Reuse-one-LESS 的质量为 34.24，低于 RDS+ 的 34.44，说明直接复用既有子集可能损失目标匹配性。PCU-Select 的复用对象是公共表征和评分器，目标配置仍保留各自子集。')
equation('cost_general')
addp(wc,'式（6）表示累计选择成本，P 为配置数，Q 为任务数；offline 表示共享离线开销，pair 为一次配置—任务组合的评分与选数成本，data 和 rank 分别表示 LESS 的建库及任务排名成本。固定 Q=4 且目标无须额外校准时，累计成本为 144+1.51P 与 63.2P，其交点为：')
equation('break_even')
addp(wc,'因此，按上述成本模型，从第三种配置开始体现摊销收益；服务五种配置时，累计选择成本分别为 151.6 与 316.0 GPU 小时，成本比约为 2.09。若仅向已有配置增加一个任务，PCU-Select 与 LESS 分别为约 0.38 和 0.30 GPU 小时，收益主要来自跨配置复用。若每个新配置的四个任务均需 500 条标签校准，则每配置另加 4.2 GPU 小时，交点移至约 2.50。上述倍数均不代表端到端训练加速。')

heading2(wc,'3.4 实验三 未见配置上的少量标签校准')
addp(wc,'本实验检查评分器超出训练注册范围后的适用边界。评测前按结构支持程度分组：L0 接近已见配置，L1 为同一家族内的较大结构变化，L2 为未见家族。图 3 汇总各组目标在 GSM8K、HumanEval 和 MMLU 三个任务上的平均差距，纵轴为 LESS 减去 PCU-Select 的分数，数值越小越好。')
addp(wc,'Cal-200 和 Cal-500 分别使用 200 条与 500 条目标配置—任务标签，冻结共享评分器，仅拟合零初始化的线性残差头。每条标签只使用一个已适应检查点和一步更新，对应成本分别为每个配置—任务 0.42 与 1.05 GPU 小时。该设计检验少量局部监督能否修正超出注册范围后的排序偏差。')
fig(TMP/'assets/calibration.png','图 3  注册范围外配置的校准结果（差距越小越好）',10.3)
addp(wc,'注：图源为论文正文，汇总三个任务。L0 未报告校准结果，图中空缺不表示零差距。',size=9,indent=False,line=14,after=4)
addp(wc,'当前稿件中，L0 的直接评分差距为 0.34 个点；L1 由 2.08 经 Cal-200 降至 0.91，再经 Cal-500 降至 0.30。Prefix/P-Tuning 由 6.97 降至 2.17 和 1.99，增加 300 条标签仅再改善 0.18 个点。该趋势表明，同家族偏移可通过少量校准缩小差距，跨家族目标仍可能保留较大误差，需结合结构支持、校准成本和目标质量选择使用范围。')

major(wc,'4 总结').paragraph_format.space_before=Pt(30)
heading2(wc,'4.1 研究总结与不足');blocks('summary')
addp(wc,'下一阶段将整理动机分析、校准与正文实验的运行记录，使背景假设、方法设计与实验结论逐项对应。')
heading2(wc,'4.2 后续研究计划')
for b in sec['future']['blocks']:
 text=b['text'].replace('以内容摘要、模型版本、预处理、种子、候选顺序和校验值描述依赖','以内容摘要、模型与评分器版本、预处理、种子、候选顺序和校验值描述依赖').replace('首版采用单节点、任务队列和本地产物库。','首版固定评分器版本，采用单工作节点、任务队列、元数据索引和本地产物库，并提供简洁控制台，支持提交实验矩阵、查看复用计划及导出报告。')
 addp(wc,text)
addp(wc,'表 3  后续八周研究计划（2026 年）',bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=4,after=4,keep=True)
schedule=method['schedule'];smalltable(wc,schedule['headers'],schedule['rows'],[2.6,6.6,6.0])
addp(wc,'各阶段以可运行产物和可核验记录为依据，根据实验资源及导师指导调整。',size=9,indent=False,line=14,after=3)

blankcell(ac);addp(ac,'与课题相关的学术成果情况：',bold=True,indent=False,line=17,after=4)
addp(ac,'已形成英文论文稿《PCU-Select: Amortizing Target-Specific Data Selection Across PEFT Configurations》，完成方法框架和现阶段实验结果的组织，正在完善原始实验记录核验及论文论证。已形成配套特征提取、PEFT 配置注册、条件评分、聚类配额选数及训练评测脚本基础，为后续系统研究提供可复用模块。')

# Preserve table grids but remove old row exceptions that trigger artificial blank pages.
for row in main.rows:
 for e in list(row._tr):
  if e.tag==qn('w:tblPrEx'):row._tr.remove(e)
 pr=row._tr.get_or_add_trPr()
 for e in pr.findall(qn('w:cantSplit')):pr.remove(e)
 cs=OxmlElement('w:cantSplit');cs.set(qn('w:val'),'0');pr.append(cs)
 for c in row.cells:
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
  for p in c.paragraphs[:2]:p.paragraph_format.keep_with_next=False

# Let the printed author declaration follow the achievements on the same page.
for p in doc.paragraphs:
 if p.text.startswith('二、中期考核作者承诺'):p.paragraph_format.page_break_before=False

# Update document identity only; cover and school review furniture are unchanged.
doc.core_properties.title='研究生学位论文中期考核表'
doc.core_properties.author='林肯达';doc.core_properties.last_modified_by='林肯达'
normalize_body_math_and_properties()
inter=TMP/'authored.docx';doc.save(inter)
allowed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','word/settings.xml'}
with ZipFile(BASE) as z:original={n:z.read(n) for n in z.namelist()}
with ZipFile(inter) as z:new={n:z.read(n) for n in z.namelist()}
added=set(new)-set(original)
assert all(n.startswith('word/media/') for n in added),added
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
 for n,b in original.items():z.writestr(n,new[n] if n in allowed else b)
 for n in added:z.writestr(n,new[n])
with ZipFile(OUT) as z:
 for n in original:
  if n not in allowed:assert z.read(n)==original[n],n
 tree=etree.fromstring(z.read('word/document.xml'))
 W={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
 assert len(tree.xpath('//m:oMathPara',namespaces=W))==eqcount
 assert len(tree.xpath('//w:drawing',namespaces=W))==4
 # Existing cover and both review tables are exact element clones.
 src=etree.fromstring(original['word/document.xml'])
 for a,b in zip(src.xpath('//w:body/w:tbl',namespaces=W)[-2:],tree.xpath('//w:body/w:tbl',namespaces=W)[-2:]):assert etree.tostring(a)==etree.tostring(b)
manifest={'base':str(BASE),'base_sha256':hashlib.sha256(BASE.read_bytes()).hexdigest(),'output':str(OUT),'equations':eqcount,'figures':3,'research_tables':3,'modified_parts':sorted(allowed),'added_parts':sorted(added),'preserved_parts':sorted(set(original)-allowed)}
(TMP/'package_validation.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(OUT);print('Native equations:',eqcount,'Original figures: 3, editable research tables: 3')
