from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree
import json,ast,hashlib
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path('/Users/bytedance/code/data-selection');TMP=ROOT/'tmp/midterm_assessment_paper'
BASE=ROOT/'output/midterm_assessment/林肯达_研究生学位论文中期考核表.docx'
OUT=ROOT/'output/midterm_assessment/林肯达_中期考核表_论文式图文版.docx'
doc=Document(BASE)
# Reuse the existing reference-derived typography and table helpers.
module=ast.parse((ROOT/'tmp/midterm_assessment/build_form.py').read_text())
for n in module.body:
 if isinstance(n,ast.FunctionDef):exec(compile(ast.Module(body=[n],type_ignores=[]),'reference_helpers','exec'))
main=doc.tables[1];ic=main.cell(2,0);wc=main.cell(3,0);ac=main.cell(4,0)
method=json.loads((TMP/'method_content.json').read_text());sec={s['key']:s for s in method['sections']}
review=json.loads((TMP/'figure_review.json').read_text())

# Native Office Math constructors; all displayed equations remain editable.
def E(tag,*children,**attrs):
 el=OxmlElement('m:'+tag)
 for k,v in attrs.items():el.set(qn('m:'+k),str(v))
 for c in children:
  if isinstance(c,list):
   for cc in c:el.append(cc)
  else:el.append(c)
 return el

def R(text,normal=False):
 r=E('r');pr=E('rPr',E('sty',val='p' if normal else 'i'));r.append(pr)
 w=OxmlElement('w:rPr');f=OxmlElement('w:rFonts');f.set(qn('w:ascii'),'Cambria Math');f.set(qn('w:hAnsi'),'Cambria Math');w.append(f)
 sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');w.append(sz);r.append(w)
 t=E('t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);return r

def seq(x):return x if isinstance(x,list) else [x if not isinstance(x,str) else R(x)]
def sub(b,s):return E('sSub',E('e',*seq(b)),E('sub',*seq(s)))
def sup(b,s):return E('sSup',E('e',*seq(b)),E('sup',*seq(s)))
def ss(b,s,u):return E('sSubSup',E('e',*seq(b)),E('sub',*seq(s)),E('sup',*seq(u)))
def frac(n,d):return E('f',E('num',*seq(n)),E('den',*seq(d)))
def group(xs,left='(',right=')'):
 return E('d',E('dPr',E('begChr',val=left),E('endChr',val=right)),E('e',*seq(xs)))
def nsum(low,term):
 return E('nary',E('naryPr',E('chr',val='∑'),E('limLoc',val='subSup'),E('supHide',val='1')),E('sub',*seq(low)),E('sup'),E('e',*seq(term)))
def hat(x):return E('acc',E('accPr',E('chr',val='̂')),E('e',*seq(x)))
def tilde(x):return E('acc',E('accPr',E('chr',val='̃')),E('e',*seq(x)))

def eqnodes(key):
 if key=='selection_objective':
  arg=E('limLow',E('e',R('arg max',True)),E('lim',R('S ⊆ D, |S| ≤ B')))
  return [sup('S','*'),R(' = '),arg,R(' '),R('Perf',True),group([R('Tune',True),group([sub('M','0'),R(', p, S')]),R(', '),ss('V','t',R('test',True))])]
 if key=='proxy_label':
  return [sup('u',R('lo',True)),R('(x,p,t) = '),nsum('ω∈Ω',[ss(tilde('w'),'p','ω'),R(' cos',True),group([ss('g','x','ω'),R(', '),ss('g','t','ω')])])]
 if key=='short_update':
  v=sub('V','t');theta=sub('θ','a')
  return [R('Δ(x,p,t,a,h) = '),sub('L',deepcopy(v)),group(deepcopy(theta)),R(' − '),sub('L',deepcopy(v)),group([ss(R('Adapt',True),'p','h'),group([R('B(x); '),deepcopy(theta)])])]
 if key=='conservative_score':
  args=[R('(x,'),sup('p','*'),R(','),sup('t','*'),R(')')]
  return [R('q(x) = '),sub(hat('μ'),'φ'),*deepcopy(args),R(' − 0.2'),sub(hat('σ'),'φ'),*deepcopy(args)]
 if key=='cluster_quota':
  def val(i):return [sup(group(ss('v',i,'+')),'0.6'),sup([R('|'),sub('C',i),R('|')],'0.4')]
  return [sub('b','k'),R(' = '),R('round',True),group([R('B '),frac(val('k'),nsum('j',val('j')))],'[',']'),R(',   '),ss('v','k','+'),R(' = '),R('max',True),group([sub('v','k'),R(',0')])]
 if key=='cost_general':
  lines=[
   [sub('C',R('PCU',True)),R('(P,Q) = '),sub('C',R('offline',True)),R(' + PQ '),ss('C',R('PCU',True),R('pair',True))],
   [sub('C',R('LESS',True)),R('(P,Q) = P '),ss('C',R('LESS',True),R('data',True)),R(' + PQ '),ss('C',R('LESS',True),R('rank',True))]]
  return [E('eqArr',*[E('e',*a) for a in lines])]
 if key=='break_even':
  return [sup('P','*'),R(' = '),frac(R('144'),R('63.2 − 1.51')),R(' ≈ 2.33')]
 raise KeyError(key)

eqcount=0

def equation(key):
 global eqcount
 eqcount+=1;p=addp(wc,'',indent=False,line=28,before=4,after=6)
 p.paragraph_format.line_spacing=1.2
 p.paragraph_format.tab_stops.add_tab_stop(Cm(7.5),WD_TAB_ALIGNMENT.CENTER)
 p.paragraph_format.tab_stops.add_tab_stop(Cm(15.2),WD_TAB_ALIGNMENT.RIGHT)
 font(p.add_run('\t'),10.5)
 math=E('oMath',*eqnodes(key));p._p.append(math)
 font(p.add_run('\t('+str(eqcount)+')'),10.5)
 p.paragraph_format.keep_together=True
 return p

def heading2(c,text):return addp(c,text,bold=True,indent=False,before=7,after=4,line=17,keep=True)
def major(c,text):return addp(c,text,bold=True,indent=False,before=9,after=5,line=18,keep=True)

def fig(path,caption,width):
 # An indivisible single-row component keeps the image and its caption together.
 t=wc.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 t.columns[0].width=Cm(15.2);t.cell(0,0).width=Cm(15.2)
 margins(t,30,0,30,0)
 b=OxmlElement('w:tblBorders');t._tbl.tblPr.append(b)
 for name in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+name);el.set(qn('w:val'),'nil');b.append(el)
 cs=OxmlElement('w:cantSplit');t.rows[0]._tr.get_or_add_trPr().append(cs)
 c=t.cell(0,0);p=setcell(c,'',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=16,after=3)
 p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
 pic=p.add_run().add_picture(str(path),width=Cm(width))
 pic._inline.docPr.set('descr',caption)
 para(c.add_paragraph(),caption,size=9.5,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,line=14,after=0)
 para(wc.paragraphs[-1],'',size=1,line=2,indent=False,after=0)
 return t

# The official form labels remain; the content now follows a paper's logic.
blankcell(ic)
addp(ic,'课题研究的创新点：',bold=True,indent=False,line=17,after=4)
major(ic,'1 背景及动机')
for s in [
'大语言模型的指令微调能够提升模型对下游任务的适应能力，但大规模候选数据和反复试验会带来较高的训练成本。参数高效微调（PEFT）通过限制可训练参数降低资源消耗，数据选择通过保留更有价值的样本减少数据开销。两类方法共同作用于训练过程：不同 PEFT 配置的更新位置、算子及容量不同，同一样本对目标任务的训练价值也可能随之变化。因而，数据选择需要同时考虑“训练什么任务”和“以何种方式更新模型”。',
'现有方法难以同时兼顾目标针对性和多配置复用。基于语义相似度或共享梯度的选择方式可以复用公共计算，但未充分表达具体 PEFT 配置的差异。逐配置 LESS 为每种配置构建专属梯度数据存储，再针对各任务评分，具有较强的目标针对性，却需要在新增配置时重复支付建库成本。直接沿用已有配置选出的子集，又可能损失与新配置的匹配性。本文由此研究可复用的条件效用评分：复用样本信号和评分器，并为不同配置与任务生成各自的训练子集。']:
 addp(ic,s)
major(ic,'2 创新点')
for s in [
'（1）面向任务与参数配置的条件效用建模。将样本、PEFT 配置与任务共同作为效用预测条件，利用共享干预位点建立可比较的信号空间，并以结构化编码区分更新位置、算子、容量和训练配方，使公共信号复用与配置差异表达相结合。',
'（2）多保真监督驱动的统一评分器。以可缓存的位点梯度代理提供低成本排序监督，再通过少量短更新观测校正代理偏差；借助条件调制和交互建模学习样本效用与不确定性，减少重复构建目标专属选择器的成本。',
'（3）兼顾效用与覆盖的目标子集构建。将保守条件评分与语义簇规模共同用于预算分配，缓解全局高分样本过于集中的问题；通过区分一次性离线成本和新增配置成本，分析多配置数据选择的摊销收益。']:
 addp(ic,s)

blankcell(wc);addp(wc,'开题以来研究工作总结：',bold=True,indent=False,line=17,after=4)

def blocks(section):
 for b in sec[section]['blocks']:
  if b['type']=='paragraph':addp(wc,b['text'])
  elif b['type']=='equation':equation(b['id'])
  elif b['type']=='figure_reference':
   fig(ROOT/b['asset'],'图 1  PCU-Select 总体架构：共享信号提取、多保真效用学习与目标条件化选择',14.2)

major(wc,'1 问题定义');blocks('problem')
major(wc,'2 方法')
heading2(wc,'2.1 总体框架');blocks('framework')
heading2(wc,'2.2 共享干预位点与结构化表征');blocks('representations')
heading2(wc,'2.3 多保真效用监督');blocks('supervision')
heading2(wc,'2.4 条件评分与聚类配额选择');blocks('selection')
addp(wc,'评分器在已见配置及结构接近的目标上直接使用；对同一家族内较大偏移或未见家族，存在兼容标签时进行小规模校准，否则采用逐目标选择方法。该支持策略将可复用评分与适用边界结合，避免将固定骨干上的共享表征直接解释为任意配置、任意骨干均可迁移。')

major(wc,'3 实验')
heading2(wc,'3.1 实验设置')
for s in [
'实验围绕下游任务质量、关键模块作用和多配置成本三方面展开。采用 Llama-2-7B 骨干，从 30 万条混合指令候选数据中选择 10%，即 3 万条样本。任务覆盖 GSM8K、HumanEval、MMLU 与 TyDiQA，分别使用精确匹配率（EM）、Pass@1、准确率（Acc）和 F1。每个任务摘要含 32 个训练或开发侧示例；HumanEval 无训练划分，使用独立代码数据构建摘要，并与最终评测数据分离。',
'主实验包括 AD-b64、IA3-attnmlp、L-r16-qkvo、L-r8-mlp 与 L-r8-qv 五种已见配置，覆盖 Adapter、IA3 和不同 LoRA 更新位置及容量。对比方法为随机选择 Random、表征相似度方法 RDS+、共享梯度方法 Influence 及逐配置 LESS。各方法采用一致的骨干、PEFT 配方和评测协议，下游训练设置为 1000 步 AdamW、全局批量 128、最大序列长度 1024。成本使用 GPU 小时统计，即墙钟时间乘以 GPU 数量。',
'以下结果汇总论文实验。表中保留均值，描述性统计不用于独立样本显著性结论。']:
 addp(wc,s)

heading2(wc,'3.2 主实验结果')
addp(wc,'表 1 汇总了各任务在五种 PEFT 配置上的平均结果。PCU-Select 的四任务宏平均为 35.18，逐配置 LESS 为 35.14，两者差值为 0.04 个指标点，整体质量接近；相对 RDS+ 和 Influence 分别增加 0.74 和 0.50 个点。')
addp(wc,'表 1  四任务主实验结果',bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=4,after=4,keep=True)
rows=[]
for r in review['main_results_table']['rows']:
 rows.append([r['method']]+[f"{r['tasks'][k]:.2f}" for k in ['GSM8K','HumanEval','MMLU','TyDiQA']]+[f"{r['macro_average']:.2f}"])
t=smalltable(wc,['方法','GSM8K\nEM','HumanEval\nPass@1','MMLU\nAcc','TyDiQA\nF1','宏平均'],rows,[3.0,2.4,2.8,2.3,2.3,2.4])
for c in t.rows[-1].cells:
 for p in c.paragraphs:
  for r in p.runs:r.bold=True
addp(wc,'注：实验结果汇总。各任务列为五种已见配置的原生百分制指标均值，宏平均为四任务等权汇总；先按未舍入数据计算，再保留两位小数。',size=9,indent=False,line=14,after=5)
addp(wc,'分任务看，PCU-Select 相对 LESS 在 GSM8K 和 MMLU 上分别增加 0.29 和 0.67 个点，在 HumanEval 和 TyDiQA 上分别减少 0.04 和 0.77 个点。平均质量接近并不表示所有任务占优，后续需重点分析 TyDiQA 的差距。')

heading2(wc,'3.3 消融实验')
addp(wc,'为分析各模块的作用，消融实验采用 GSM8K、HumanEval 与两个代表性 PEFT 配置，预算仍为 10%。表 2 的完整方法宏平均为 19.72，其统计范围与表 1 不同。移除 PEFT 编码产生最大的性能下降，表明配置条件是可复用评分器的重要信息；全局 Top-k 和仅代理监督也出现下降，支持在当前设置下保留语义覆盖与短更新校正。')
addp(wc,'表 2  主要模块消融结果',bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=4,after=4,keep=True)
arows=[]
for name,value,drop in review['ablation_table_or_optional_redraw']['selected_rows']:
 if name not in ['完整PCU-Select','移除PEFT编码','移除任务摘要','仅用梯度代理标签','改用全局Top-k']:continue
 arows.append([name,f'{value:.2f}','—' if drop==0 else f'{drop:.2f}'])
smalltable(wc,['设置','宏平均','相对完整方法下降'],arows,[7.4,3.4,4.4])
addp(wc,'注：实验结果汇总。仅汇总两个任务与两个代表配置，下降量单位为原生指标点。',size=9,indent=False,line=14,after=5)

heading2(wc,'3.4 多配置成本与摊销分析')
addp(wc,'设 P 为待服务的 PEFT 配置数，Q 为任务数。LESS 为每种配置构建一次梯度数据存储，再对各任务排名；PCU-Select 将表征提取、标签构建和评分器训练集中在共享离线阶段，之后为每个配置—任务组合执行评分与选数。因此，累计选择成本可表示为：')
equation('cost_general')
addp(wc,'其中，offline 为公共离线开销，pair 为一次目标组合的评分与选数成本，data 和 rank 分别表示 LESS 的建库与任务排名成本。固定 Q=4 且配置无须额外校准时，当前稿件给出 PCU-Select 成本为 144+1.51P，LESS 为 63.2P，单位均为 GPU 小时。图 2 按这些分项成本绘制累计曲线；交点满足：')
equation('break_even')
fig(TMP/'assets/cost_curve.png','图 2  随 PEFT 配置数增加的累计选择成本（每配置四任务）',11.7)
addp(wc,'注：图据论文分项成本计算。包含一次性离线开销，排除共同的下游训练；横轴为配置数，图中 T*=2.3 是 P*≈2.33 的舍入显示。',size=9,indent=False,line=14,after=5)
addp(wc,'因此，从第三个完整配置开始体现摊销收益。五配置的累计选择成本约为 151.6 和 316.0 GPU 小时，成本比约为 2.09；63.2/1.51≈42 仅表示新增四任务配置的边际选择成本比。新增任务与新增配置的成本结构不同，不能将上述比值解释为端到端训练加速。需要额外校准的目标还应计入校准标签成本。')

major(wc,'4 总结')
heading2(wc,'4.1 研究总结与不足');blocks('summary')
heading2(wc,'4.2 后续研究计划')
for b in sec['future']['blocks']:
 text=b['text'].replace('以内容摘要、模型版本、预处理、种子、候选顺序和校验值描述依赖','以内容摘要、模型与评分器版本、预处理、种子、候选顺序和校验值描述依赖').replace('首版采用单节点、任务队列和本地产物库。','首版固定评分器版本，采用单工作节点、任务队列、元数据索引和本地产物库，并提供简洁控制台，支持提交实验矩阵、查看复用计划及导出报告。')
 addp(wc,text)
addp(wc,'表 3  后续八周研究计划（2026 年）',bold=True,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=4,after=4,keep=True)
schedule=method['schedule'];smalltable(wc,schedule['headers'],schedule['rows'],[2.6,6.6,6.0])
addp(wc,'各阶段以可运行产物和可核验记录为依据，根据实验资源及导师指导调整。',size=9,indent=False,line=14,after=3)

blankcell(ac);addp(ac,'与课题相关的学术成果情况：',bold=True,indent=False,line=17,after=4)
addp(ac,'已形成英文论文稿《PCU-Select: Amortizing Target-Specific Data Selection Across PEFT Configurations》，完成方法框架和现阶段实验结果的组织，正在完善原始实验记录核验及论文论证。已形成配套特征提取、PEFT 配置注册、条件评分、聚类配额选数及训练评测脚本基础，为后续系统研究提供可复用模块。')

# Preserve table grids but remove old row exceptions that trigger artificial blank pages.
for row in main.rows:
 for e in list(row._tr):
  if e.tag==qn('w:tblPrEx'):row._tr.remove(e)
 pr=row._tr.get_or_add_trPr()
 for e in pr.findall(qn('w:cantSplit')):pr.remove(e)
 cs=OxmlElement('w:cantSplit');cs.set(qn('w:val'),'0');pr.append(cs)
 for c in row.cells:
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
  for p in c.paragraphs[:2]:p.paragraph_format.keep_with_next=False

# Let the printed author declaration follow the achievements on the same page.
for p in doc.paragraphs:
 if p.text.startswith('二、中期考核作者承诺'):p.paragraph_format.page_break_before=False

# Update document identity only; cover and school review furniture are unchanged.
doc.core_properties.title='研究生学位论文中期考核表'
doc.core_properties.author='林肯达';doc.core_properties.last_modified_by='林肯达'
inter=TMP/'authored.docx';doc.save(inter)
allowed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml'}
with ZipFile(BASE) as z:original={n:z.read(n) for n in z.namelist()}
with ZipFile(inter) as z:new={n:z.read(n) for n in z.namelist()}
added=set(new)-set(original)
assert all(n.startswith('word/media/') for n in added),added
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
 for n,b in original.items():z.writestr(n,new[n] if n in allowed else b)
 for n in added:z.writestr(n,new[n])
with ZipFile(OUT) as z:
 for n in original:
  if n not in allowed:assert z.read(n)==original[n],n
 tree=etree.fromstring(z.read('word/document.xml'))
 W={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
 assert len(tree.xpath('//m:oMath',namespaces=W))==eqcount
 assert len(tree.xpath('//w:drawing',namespaces=W))==3
 # Existing cover and both review tables are exact element clones.
 src=etree.fromstring(original['word/document.xml'])
 for a,b in zip(src.xpath('//w:body/w:tbl',namespaces=W)[-2:],tree.xpath('//w:body/w:tbl',namespaces=W)[-2:]):assert etree.tostring(a)==etree.tostring(b)
manifest={'base':str(BASE),'base_sha256':hashlib.sha256(BASE.read_bytes()).hexdigest(),'output':str(OUT),'equations':eqcount,'figures':2,'research_tables':3,'modified_parts':sorted(allowed),'added_parts':sorted(added),'preserved_parts':sorted(set(original)-allowed)}
(TMP/'package_validation.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(OUT);print('Native equations:',eqcount,'Original figures: 2, editable research tables: 3')
