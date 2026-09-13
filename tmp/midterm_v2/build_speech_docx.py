import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/bytedance/code/data-selection')
DATA = json.loads((ROOT/'tmp/midterm_v2/content.json').read_text())
OUT = ROOT/'output/midterm_v2'
OUT.mkdir(parents=True, exist_ok=True)
STEM = '林肯达_中期答辩讲稿_论文重点版'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.68)
sec.bottom_margin = Inches(.64)
sec.left_margin = Inches(.80)
sec.right_margin = Inches(.80)
sec.header_distance = Inches(.20)
sec.footer_distance = Inches(.29)

for name, size in [('Normal',12),('Title',25),('Subtitle',13.5),('Heading 1',15),('Heading 2',13.1),('Heading 3',12)]:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    st.font.bold = name.startswith('Heading') or name == 'Title'
    st.font.italic = False
    rp = st.element.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rp.append(rf)
    for key in list(rf.attrib):
        if 'theme' in key.lower(): del rf.attrib[key]
    for key,value in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia','Arial Unicode MS'),('cs','Arial')]:
        rf.set(qn('w:'+key),value)
    pf = st.paragraph_format
    pf.line_spacing = Pt(18 if name=='Normal' else size*1.3)
    pf.space_after = Pt(5)
    pf.widow_control = True
    if name.startswith('Heading'):
        pf.space_before = Pt(9)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
    if name=='Title': pf.space_after=Pt(12)

for rf in doc.styles.element.iter(qn('w:rFonts')):
    for key in list(rf.attrib):
        if 'theme' in key.lower(): del rf.attrib[key]
    rf.set(qn('w:eastAsia'),'Arial Unicode MS')
for color in doc.styles.element.iter(qn('w:color')):
    for key in list(color.attrib):
        if 'theme' in key.lower(): del color.attrib[key]
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

doc.core_properties.title = '林肯达中期答辩讲稿 论文重点版'
doc.core_properties.subject = DATA['title']
doc.core_properties.author = '林肯达'
doc.core_properties.keywords = '中期答辩 PCU-Select 逐页讲稿'

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.line_spacing=Pt(11)
for part in ['第 ', None, ' 页']:
    if part is None:
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    else:
        run=footer.add_run(part);run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('626262')

def clean(t):
    return re.sub(r'\s+',' ',re.sub(r'[^\w\u4e00-\u9fff ]',' ',t)).strip()

def para(t,style=None,size=None,bold=False,color=None,after=5,keep=False):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.keep_together=True
    p.paragraph_format.keep_with_next=keep
    r=p.add_run(t)
    r.bold=bold if bold else None
    if size:r.font.size=Pt(size)
    if color:r.font.color.rgb=RGBColor.from_string(color)
    return p

def note(label,t,after=5,keep=False,color='555555'):
    p=doc.add_paragraph()
    p.paragraph_format.line_spacing=Pt(14)
    p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.keep_together=True
    p.paragraph_format.keep_with_next=keep
    r=p.add_run(label+'  ');r.bold=True;r.font.size=Pt(10)
    r.font.color.rgb=RGBColor.from_string(color)
    r=p.add_run(t);r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string(color)
    return p

def heading(t,level=1):
    return doc.add_heading(clean(t),level=level)

def mmss(s): return f'{s//60:02d}:{s%60:02d}'

def split_speech(text):
    if text.startswith('这里比较三类方案。'):
        end=text.index('但不区分 PEFT 配置。')+len('但不区分 PEFT 配置。')
        return [text[:end],text[end:]]
    if len(text)<220: return [text]
    ends=[m.end() for m in re.finditer('。',text)]
    valid=[e for e in ends if 70<e<len(text)-65]
    if not valid:return [text]
    cut=min(valid,key=lambda e:abs(e-len(text)/2))
    return [text[:cut],text[cut:]]

times={}
cumulative=0
for s in DATA['slides'][:28]:
    times[s['n']]=(cumulative,cumulative+s['seconds'])
    cumulative+=s['seconds']

def speak(s):
    # Preserve the presentation's final titles verbatim for slide-by-slide use.
    doc.add_heading(f"第 {s['n']:02d} 页  {s['title']}",level=2)
    if s['n']<=28:
        a,b=times[s['n']]
        note('时间',f"建议 {s['seconds']} 秒    累计 {mmss(a)} 至 {mmss(b)}",after=5,keep=True)
    else:
        note('使用', '备用页 按问题选讲 不计入主讲时间',after=5,keep=True)
    for text in split_speech(s['script']):
        para(text,after=4,keep=True)
    if s['n']!=28:
        note('转场',s['transition'],after=5,keep=True,color='222222')
    note('讲述提示',s['cue'],after=10,keep=False)

para('林肯达中期答辩讲稿','Title')
para('论文重点版','Subtitle',after=13)
para(DATA['title'],size=13.4,bold=True,after=16)
para('本讲稿对应 33 页中期答辩 PPT。第 01 至 28 页为主讲，默认约 20 分钟；第 29 至 33 页为备用，按提问使用。主体重点讲清 PCU-Select 的研究问题、方法设计及稿件结果，随后介绍承接算法的系统方案与后续计划。',after=12)
heading('时间安排')
for a,b,label in [(1,2,'开场与目录'),(3,7,'研究背景与动机'),(8,14,'PCU Select 方法'),(15,21,'实验结果与工作一小结'),(22,28,'系统计划与总结')]:
    total=sum(x['seconds'] for x in DATA['slides'][a-1:b])
    note(f'第 {a:02d} 至 {b:02d} 页',f'{label}    {total//60} 分 {total%60:02d} 秒',after=7)
heading('怎样使用讲稿')
para('正文和“转场”可直接口述；灰色的时间与讲述提示用于排练，不需要朗读。先完成一次连续计时，再根据自己的语速调整停顿。方法页优先讲清直觉和流程，数字页先说明指标口径，再解释比较结论。',after=8)
para('全文默认按 20 分钟准备。工作一约 15 分 31 秒，工作二与计划及结束约 3 分 54 秒。备用页和 12 个常见问答用于应对追问，不占主讲时间。',after=8)
heading('术语与材料说明')
note('术语','PEFT 指参数高效微调；task sketch 指任务样本集；FiLM 指按条件缩放和平移特征；GPU 小时指墙钟时间乘 GPU 数量。',after=8)
para('数字页为实验结果汇总；具体统计口径见备用第 32 页。系统功能与八周安排属于后续计划。',size=10.5,after=4)

groups=[
    (1,3,'逐页讲稿 01 至 03'),
    (4,6,'逐页讲稿 04 至 06'),
    (7,8,'逐页讲稿 07 至 08'),
    (9,10,'逐页讲稿 09 至 10'),
    (11,12,'逐页讲稿 11 至 12'),
    (13,14,'逐页讲稿 13 至 14'),
    (15,17,'逐页讲稿 15 至 17'),
    (18,19,'逐页讲稿 18 至 19'),
    (20,22,'逐页讲稿 20 至 22'),
    (23,25,'逐页讲稿 23 至 25'),
    (26,28,'逐页讲稿 26 至 28'),
    (29,30,'备用讲稿 29 至 30'),
    (31,33,'备用讲稿 31 至 33'),
]
for a,b,label in groups:
    doc.add_page_break(); heading(label)
    for s in DATA['slides'][a-1:b]:speak(s)

for start in range(0,12,3):
    doc.add_page_break();heading(f'常见问答 {start+1:02d} 至 {start+3:02d}')
    for i,q in enumerate(DATA['qa'][start:start+3],start+1):
        question='与 LESS 的均值接近能否说明显著领先' if i==6 else q['question']
        heading(f"{i:02d}  {question}",2)
        for t in split_speech(q['answer']):para(t,keep=True)
        note('对应 PPT',str(q['slides']),after=15)

doc.add_page_break();heading('内容来源与核验索引')
para('以下索引用于排练前核对数字、公式和计划口径，答辩时不需要朗读。来源依次对应题目与规划、论文方法、稿件结果、工程基础和数据核验状态。',after=13)
source_notes=[
    ('开题材料与后续方案','开题答辩 PPT 提供原总题目和汇报人信息。第二工作点采用用户提供的系统目标、两项工程机制、24 次训练示例和首版范围。'),
    ('论文问题定义与方法章节','对应第 03 至 14 页及备用第 33 页。用于核对条件效用、共享位点、三类表示、梯度代理、短更新标签、条件评分和聚类配额。'),
    ('论文实验章节与主结果表','对应第 06 页、第 15 至 17 页、第 20 页和备用第 29 页。包括主表、逐任务汇总及跨配置复用比较，主讲仅使用已审阅的均值口径。'),
    ('论文消融表与成本模型','对应第 18 至 19 页和备用第 30 页。消融是两任务和两代表配置的平均；成本比较以 GPU 小时计，区分前期、边际、校准和下游训练。'),
    ('仓库算法与实验脚本','对应第 21 至 25 页及备用第 31 页。已有基础包括特征缓存、任务表示、PEFT 注册、条件评分、聚类配额和实验脚本；分层依赖、公共节点与阶段恢复属于计划新增。'),
    ('结果审计与修订记录','对应备用第 32 页。主表按训练种子汇总，LN-Tuning 对应未见家族的校准实验；复现实验时需保留对应运行记录和计时日志。'),
]
for name,text in source_notes:
    heading(name,2);para(text,after=10)
para('最终验收关注可运行系统、可追溯实验产物、两项机制的定量报告，以及由真实实验记录支撑的论文结果。',after=4)

for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
out=OUT/(STEM+'.docx')
doc.save(out)

md=['# 林肯达中期答辩讲稿 论文重点版','',DATA['title'],'',
    '主讲 28 页，默认 20 分钟；备用 5 页按提问选用。工作一约 15 分 31 秒，系统与计划及结束约 3 分 54 秒，开场目录 35 秒。','',
    '正文与转场可直接口述；时间和讲述提示用于排练，不需要朗读。数字为实验结果汇总；系统功能与八周安排为计划。','']
for s in DATA['slides']:
    md += [f"## 第 {s['n']:02d} 页 {s['title']}",'']
    if s['n']<=28:
        a,b=times[s['n']];md += [f"建议 {s['seconds']} 秒；累计 {mmss(a)} 至 {mmss(b)}。",'']
    else:md += ['备用页，按问题选讲，不计入主讲时间。','']
    for text in split_speech(s['script']):md += [text,'']
    if s['n']!=28:md += [f"转场：{s['transition']}",'']
    md += [f"讲述提示：{s['cue']}",'']
md += ['## 常见问答','']
for i,q in enumerate(DATA['qa'],1):
    md += [f"### {i:02d} {q['question']}",'',q['answer'],'',f"对应 PPT：{q['slides']}",'']
md += ['## 内容来源与核验索引','']
for name,text in source_notes:md += [f'### {name}','',text,'']
(OUT/(STEM+'.md')).write_text('\n'.join(md),encoding='utf-8')
print(out)
print(OUT/(STEM+'.md'))
print('Expected document pages:',1+len(groups)+4+1)
