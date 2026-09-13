import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/Users/bytedance/code/data-selection')
DATA=json.loads((ROOT/'tmp/defense_build/defense_content.json').read_text())
OUT=ROOT/'output/defense'
OUT.mkdir(parents=True,exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.85)
sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.15)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.8)

def set_fonts(style,east='Arial Unicode MS',latin='Arial',size=12,bold=False):
    style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    el=style.element.get_or_add_rPr()
    f=el.find(qn('w:rFonts'))
    if f is None:f=OxmlElement('w:rFonts');el.append(f)
    for k in list(f.attrib):
        if 'theme' in k.lower():del f.attrib[k]
    for k,v in [('ascii',latin),('hAnsi',latin),('eastAsia',east),('cs',latin)]:f.set(qn('w:'+k),v)
    for b in list(style.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    for color in style.element.iter(qn('w:color')):
        for k in list(color.attrib):
            if 'theme' in k.lower():del color.attrib[k]

set_fonts(doc.styles['Normal'])
normal=doc.styles['Normal'].paragraph_format
normal.line_spacing=Pt(18);normal.space_after=Pt(7)
set_fonts(doc.styles['Title'],east='Arial Unicode MS',size=26,bold=True)
doc.styles['Title'].paragraph_format.space_after=Pt(17)
doc.styles['Title'].paragraph_format.line_spacing=Pt(34)
for name,size in [('Heading 1',18),('Heading 2',14.5),('Heading 3',12.5)]:
    set_fonts(doc.styles[name],east='Arial Unicode MS',size=size,bold=True)
    doc.styles[name].paragraph_format.space_before=Pt(9)
    doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles[name].paragraph_format.line_spacing=Pt(size*1.3)
set_fonts(doc.styles['Subtitle'],east='Arial Unicode MS',size=13)
doc.styles['Subtitle'].font.italic=False
doc.styles['Subtitle'].paragraph_format.line_spacing=Pt(20)

doc.core_properties.title='PCU Select 论文答辩讲稿'
doc.core_properties.subject='跨 PEFT 配置复用的数据选择'
doc.core_properties.author=''
doc.core_properties.keywords='PCU-Select, PEFT, 答辩, 逐页讲稿'

foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
for r in foot.runs:r.font.size=Pt(9)

def para(text,style=None,bold=False,size=None,color=None):
    p=doc.add_paragraph(style=style)
    r=p.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    if color:r.font.color.rgb=RGBColor.from_string(color)
    return p

def note(label,text):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.line_spacing=Pt(14)
    r=p.add_run(label+' ');r.bold=True;r.font.size=Pt(10.5)
    r=p.add_run(text);r.font.size=Pt(10.5)
    r.font.color.rgb=RGBColor.from_string('4D5660')
    return p

def newpage(title):
    doc.add_page_break()
    doc.add_heading(title,level=1)

def speak(slide,backup=False):
    title=slide['title'].replace('备用  ','')
    doc.add_heading(f"第 {slide['n']:02d} 页  {title}",level=2)
    if not backup:
        p=para(f"建议 {slide['seconds']} 秒",size=10,color='4D5660')
        p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.keep_with_next=True
    speech=para(slide['script'])
    speech.paragraph_format.keep_together=True
    speech.paragraph_format.keep_with_next=True
    if slide['transition']:
        transition=note('转场',slide['transition'])
        transition.paragraph_format.keep_with_next=True
    note('提示',slide['cue'])

doc.add_paragraph('PCU Select 论文答辩讲稿',style='Title')
doc.add_paragraph('跨 PEFT 配置复用的数据选择',style='Subtitle')
para('本讲稿对应 18 页正文和 4 页备用 PPT。正文按约 15 分钟设计，逐页提供可直接讲述的口语稿、转场句和提示。最后附常见追问的参考回答。')
doc.add_heading('汇报主线',level=1)
para('先解释配置变化为什么可能改变数据价值，再介绍共享位点、条件评分器和选集构建。实验部分分别讨论质量、成本与适用边界。全文的核心是：保留目标配置的信息，让昂贵的选数计算可以跨配置复用。')
doc.add_heading('时间分配',level=1)
rows=[['PPT 页码','内容','建议时间'],['01 至 05','问题与研究目标','3 分 15 秒'],['06 至 10','方法设计','4 分 35 秒'],['11 至 13','评价与质量比较','2 分 25 秒'],['14 至 16','成本 消融与迁移','2 分 40 秒'],['17 至 18','边界与总结','1 分 20 秒']]
t=doc.add_table(rows=len(rows),cols=3)
t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[Cm(3.5),Cm(9.0),Cm(4.1)]
for col,width in zip(t.columns,widths):col.width=width
for i,row in enumerate(rows):
    for j,val in enumerate(row):
        c=t.cell(i,j);c.width=widths[j];c.text=val;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr=c._tc.get_or_add_tcPr()
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'17314D' if i==0 else ('F0F3F6' if i%2==0 else 'FFFFFF'));tcPr.append(sh)
        mar=OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            el=OxmlElement('w:'+side);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');mar.append(el)
        tcPr.append(mar)
        borders=OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
        tcPr.append(borders)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(14)
            for r in p.runs:
                r.font.size=Pt(10.5);r.bold=i==0
                if i==0:r.font.color.rgb=RGBColor(255,255,255)
para('正文计划用时 14 分 15 秒，预留约 45 秒用于停顿和翻页。首次排练建议计时，方法页可以放慢，备用页只在问答时使用。',size=10.5)
doc.add_heading('实验数据的表述',level=1)
para('本讲稿按论文实验结果说明质量、成本与配置迁移表现。')

for start in range(0,18,3):
    newpage(f'逐页讲稿  {start+1:02d} 至 {start+3:02d}')
    for sl in DATA['slides'][start:start+3]:speak(sl)

newpage('备用页讲述')
for sl in DATA['slides'][18:]:speak(sl,True)

for start in range(0,12,4):
    newpage(f'常见追问  {start+1:02d} 至 {start+4:02d}')
    for i,q in enumerate(DATA['questions'][start:start+4],start+1):
        doc.add_heading(f"{i:02d}  {q['q']}",level=2)
        para(q['a'])
        note('对应 PPT',q['slides'])

for fonts in doc.styles.element.iter(qn('w:rFonts')):
    for key in list(fonts.attrib):
        if 'theme' in key.lower():del fonts.attrib[key]
    fonts.set(qn('w:eastAsia'),'Arial Unicode MS')
for borders in list(doc.styles.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
for borders in list(doc.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
doc.save(OUT/'PCU-Select_逐页答辩讲稿.docx')

md=['# PCU-Select 论文答辩讲稿','',DATA['subtitle'],'',f"正文 18 页，备用 4 页，{DATA['duration']}。",'','## 数据状态','',DATA['evidenceStatus'],'']
for sl in DATA['slides']:
    md += [f"## 第 {sl['n']:02d} 页 {sl['title']}",'',f"建议讲述：{sl['seconds']} 秒" if sl['seconds'] else '备用页，按需讲述。','',f"本页核心：{sl['takeaway']}",'',sl['script'],'']
    if sl['transition']:md += [f"转场：{sl['transition']}",'']
    md += [f"提示：{sl['cue']}",'']
md+=['## 常见追问','']
for i,q in enumerate(DATA['questions'],1):md += [f"### {i}. {q['q']}",'',q['a'],'',f"对应 PPT：{q['slides']}",'']
(OUT/'PCU-Select_逐页答辩讲稿.md').write_text('\n'.join(md))
print('Saved DOCX and Markdown speech')
