import json,re
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path('/Users/bytedance/code/data-selection')
DATA=json.loads((ROOT/'tmp/midterm_build/midterm_content.json').read_text())
OUT=ROOT/'output/midterm'
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.85);sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.15)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.8)
for name,size in [('Normal',12),('Title',26),('Subtitle',13),('Heading 1',18),('Heading 2',14.5),('Heading 3',12.5)]:
 st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=name.startswith('Heading') or name=='Title'
 rp=st.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');rp.append(rf)
 for k in list(rf.attrib):
  if 'theme' in k.lower():del rf.attrib[k]
 for k,v in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia','Arial Unicode MS'),('cs','Arial')]:rf.set(qn('w:'+k),v)
 for co in rp.iter(qn('w:color')):
  for k in list(co.attrib):
   if 'theme' in k.lower():del co.attrib[k]
 st.paragraph_format.line_spacing=Pt(size*1.35 if name!='Normal' else 18)
 st.paragraph_format.space_after=Pt(7)
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(10);st.paragraph_format.keep_with_next=True
for fonts in doc.styles.element.iter(qn('w:rFonts')):
 for key in list(fonts.attrib):
  if 'theme' in key.lower():del fonts.attrib[key]
 fonts.set(qn('w:eastAsia'),'Arial Unicode MS')
for borders in list(doc.styles.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
doc.core_properties.title='林肯达中期答辩逐页讲稿';doc.core_properties.author='林肯达'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def para(t,style=None,size=None):
 p=doc.add_paragraph(style=style);r=p.add_run(t)
 if size:r.font.size=Pt(size)
 return p

def note(label,t):
 p=doc.add_paragraph();p.paragraph_format.line_spacing=Pt(14);p.paragraph_format.space_after=Pt(6)
 r=p.add_run(label+' ');r.bold=True;r.font.size=Pt(10.5)
 r=p.add_run(t);r.font.size=Pt(10.5);r.font.color.rgb=RGBColor.from_string('52586B')
 return p

def heading(t,level=1):doc.add_heading(re.sub(r'[：:？?，,。·\-（）()]',' ',t),level=level)
def newpage(t):doc.add_page_break();heading(t)
def speak(d):
 heading(f"第 {d['n']:02d} 页  {d['title']}",2)
 if d['n']<=25:note('时间',f"建议 {d['seconds']} 秒").paragraph_format.keep_with_next=True
 speech=para(d['script']);speech.paragraph_format.keep_together=True;speech.paragraph_format.keep_with_next=True
 if d.get('transition'):note('转场',d['transition']).paragraph_format.keep_with_next=True
 note('提示',d['cue'])

para('林肯达中期答辩逐页讲稿','Title')
para('面向大模型的数据与参数联合高效微调算法研究','Subtitle')
para('本讲稿对应优化后的 27 页 PPT，其中 25 页用于主讲，2 页供问答使用。主线是介绍 PCU Select 当前进展，再说明多配置实验管理系统的两项工程机制和后续工作安排。每页提供口语讲稿、转场句与讲述提示。')
heading('使用与时间安排')
total=sum(s['seconds'] for s in DATA['slides'][:25])
para(f'正文计划用时 {total//60} 分 {total%60} 秒，按约 15 分钟准备。首次排练请计时，章节页简短过渡，算法流程与系统复用规则可以放慢。问答时按需要跳转第 26 或 27 页。')
for a,b,label in [(1,5,'课题进展概览'),(6,12,'PCU Select 方法与当前进展'),(13,20,'系统方案与工程机制'),(21,25,'后续计划与总结')]:
 t=sum(x['seconds'] for x in DATA['slides'][a-1:b]);note(f'第 {a:02d} 至 {b:02d} 页',f'{label}，约 {t//60} 分 {t%60} 秒。')
heading('汇报前需要核实的内容')
para('PPT 沿用原开题总题目和汇报人信息。后续安排为建议的相对 8 周计划，应按学院截止日期和实际实验资源调整。系统部分是拟开展工作，不能讲成已经完成。')
para('实验数据已更新为真实运行结果。')
# Group short cover/divider slides with nearby content for a practical handout.
for a,b in [(1,4),(5,7),(8,10),(11,13),(14,16),(17,19),(20,22),(23,25)]:
 newpage(f'逐页讲稿  {a:02d} 至 {b:02d}')
 for d in DATA['slides'][a-1:b]:speak(d)
newpage('备用页讲述')
for d in DATA['slides'][25:]:speak(d)
qs=DATA.get('questions',DATA.get('qa',[]))
for start in range(0,len(qs),4):
 newpage(f'常见追问  {start+1:02d} 至 {min(start+4,len(qs)):02d}')
 for i,q in enumerate(qs[start:start+4],start+1):
  heading(f"{i:02d}  {q.get('question',q.get('q',''))}",2)
  para(q.get('answer',q.get('a','')))
  if q.get('slides'):note('对应 PPT',str(q['slides']))
for borders in list(doc.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
out=OUT/'林肯达_中期答辩逐页讲稿.docx';doc.save(out)
md=['# 林肯达中期答辩逐页讲稿','','正文 25 页，备用 2 页，约 15 分钟。','']
for d in DATA['slides']:
 md += [f"## 第 {d['n']:02d} 页 {d['title']}",'',f"建议 {d['seconds']} 秒" if d['seconds'] else '备用页','',d['script'],'']
 if d.get('transition'):md += ['转场：'+d['transition'],'']
 md+=['提示：'+d['cue'],'']
md+=['## 常见追问','']
for i,q in enumerate(qs,1):md += [f"### {i}. {q.get('question',q.get('q',''))}",'',q.get('answer',q.get('a','')),'']
(OUT/'林肯达_中期答辩逐页讲稿.md').write_text('\n'.join(md))
print(out)
